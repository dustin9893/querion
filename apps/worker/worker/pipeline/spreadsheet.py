"""Read .xlsx workbooks into self-describing lines and pack them into chunks.

Each chunk is embedded on its own, so a row of values is useless without its column
names: every data row becomes ``Name: value | Name: value``. Rows are never split
across chunks (unless a single row is longer than a chunk), and every chunk carries
``[Sheet <name> › dòng a–b]`` so the citation UI can point at the Excel rows.
Spreadsheets never go through the prose chunker, which collapses whitespace and cuts
fixed windows through the middle of rows.
"""

import datetime as dt
import logging
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

from worker.pipeline.chunker import CHUNK_SIZE, ChunkResult

logger = logging.getLogger(__name__)

MAX_XML_BYTES = 16 * 1024 * 1024
MAX_CHUNKS = 2000
MAX_CELL_CHARS = 500
MAX_NAME_CHARS = 60
MAX_TITLE_CHARS = 200
PAIR_SEP = " | "

INVALID_MSG = (
    "Tệp không phải bảng tính .xlsx hợp lệ (tệp .xls cũ, tệp có mật khẩu hoặc tệp hỏng). "
    "Hãy mở bằng Excel và lưu lại dạng .xlsx."
)
EMPTY_MSG = (
    "Bảng tính không có dữ liệu đọc được. Nếu bảng chỉ gồm công thức, hãy mở bằng Excel "
    "rồi lưu lại để có giá trị."
)

_SHEET_NAME_MAX = 80
_HEAD_LABEL_MAX = 100
_MIN_BODY = 200
_WS = re.compile(r"\s+")


def _too_large_msg(xml_bytes: int) -> str:
    mb = xml_bytes / (1024 * 1024)
    limit = MAX_XML_BYTES // (1024 * 1024)
    return (
        f"Bảng tính quá lớn để lập chỉ mục (nội dung {mb:.0f} MB, tối đa {limit} MB). "
        "Hãy tách thành nhiều tệp nhỏ hơn."
    )


def _too_many_chunks_msg(n: int) -> str:
    return (
        f"Bảng tính tạo ra {n} đoạn, vượt giới hạn {MAX_CHUNKS} đoạn cho một tài liệu. "
        "Hãy tách thành nhiều tệp nhỏ hơn."
    )


def is_spreadsheet(file_path: str) -> bool:
    """True for a .xlsx file. By extension only: the API admits files by extension."""
    return Path(file_path).suffix.lower() == ".xlsx"


@dataclass
class SheetLines:
    name: str
    title: str
    lines: list[tuple[int, str, bool]] = field(default_factory=list)  # (excel_row, text, is_label)


# ---------------------------------------------------------------------------
# Reading
# ---------------------------------------------------------------------------


def _clean(text: str, limit: int, ellipsis_extra: bool = True) -> str:
    """Collapse whitespace, swap < > for fullwidth signs (the guard strips <...>), cap length."""
    text = _WS.sub(" ", text).strip().replace("<", "＜").replace(">", "＞")
    if len(text) > limit:
        cut = limit if ellipsis_extra else limit - 1
        text = text[:cut].rstrip() + "…"
    return text


def _strip_zeros(s: str) -> str:
    return s.rstrip("0").rstrip(",") if "," in s else s


def _format_number(value: float, number_format: str | None) -> str:
    section = (number_format or "General").split(";")[0]
    section = re.sub(r'"[^"]*"', "", section)
    section = re.sub(r"\[[^\]]*\]", "", section)
    m = re.search(r"\.([0#?]+)", section)
    decimals = len(m.group(1)) if m else 0

    if "%" in section:
        s = f"{value * 100:.{decimals}f}".replace(".", ",")
        return _strip_zeros(s) + "%"
    if "," in section:
        sign = "-" if value < 0 else ""
        s = f"{abs(value):,.{decimals}f}".replace(",", " ").replace(".", ",").replace(" ", ".")
        return sign + _strip_zeros(s)
    # General and anything else: no grouping, so codes (1001) and years (2026) stay as typed.
    if isinstance(value, int):
        return str(value)
    if value.is_integer():
        return str(int(value))
    return _strip_zeros(f"{value:.10f}".replace(".", ","))


def _cell_text(cell) -> str:
    """Display text of a cell as a Vietnamese reader would expect it."""
    value = cell.value
    if value is None or cell.data_type == "e":
        return ""
    if isinstance(value, bool):
        return "Có" if value else "Không"
    if isinstance(value, dt.datetime):
        text = value.strftime("%d/%m/%Y")
        if value.time() != dt.time(0, 0):
            text += value.strftime(" %H:%M")
        return text
    if isinstance(value, dt.date):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, dt.time):
        return value.strftime("%H:%M")
    if isinstance(value, dt.timedelta):
        minutes = int(value.total_seconds()) // 60
        return f"{minutes // 60}:{minutes % 60:02d}"
    if isinstance(value, (int, float)):
        return _format_number(value, cell.number_format)
    return _clean(str(value), MAX_CELL_CHARS)


def _check_zip(source) -> None:
    """Reject non-zip input and workbooks whose XML would not fit in memory, before loading."""
    try:
        with zipfile.ZipFile(source) as zf:
            xml_bytes = sum(i.file_size for i in zf.infolist() if i.filename.endswith(".xml"))
    except zipfile.BadZipFile:
        raise ValueError(INVALID_MSG) from None
    finally:
        if hasattr(source, "seek"):
            source.seek(0)
    if xml_bytes > MAX_XML_BYTES:
        raise ValueError(_too_large_msg(xml_bytes))


def _hidden_columns(ws) -> set[int]:
    from openpyxl.utils import column_index_from_string

    hidden: set[int] = set()
    for dim in ws.column_dimensions.values():
        if not dim.hidden:
            continue
        # A hidden group K:M is stored once under "K" with min=11, max=13.
        if dim.min is not None and dim.max is not None:
            hidden.update(range(dim.min, dim.max + 1))
        else:
            hidden.add(column_index_from_string(dim.index))
    return hidden


def _distinct(texts: list[str]) -> list[str]:
    return list(dict.fromkeys(t for t in texts if t))


def _read_sheet(ws) -> tuple[SheetLines, int, int]:
    """Lines of one visible worksheet, plus the number of hidden rows/columns skipped."""
    from openpyxl.cell.cell import MergedCell
    from openpyxl.utils import get_column_letter

    hidden_rows = {r for r, dim in ws.row_dimensions.items() if dim.hidden}
    hidden_cols = _hidden_columns(ws)

    # Walk the stored cells instead of iter_rows(): iter_rows() creates every cell of the
    # max_row x max_column rectangle, which explodes on a sparse sheet.
    texts: dict[tuple[int, int], str] = {}
    for (r, c), cell in ws._cells.items():
        if isinstance(cell, MergedCell):
            continue
        text = _cell_text(cell)
        if text:
            texts[r, c] = text
    max_row = max((r for r, _ in texts), default=0)
    max_col = max((c for _, c in texts), default=0)

    # Every cell of a merged range reads as its anchor. Clamp to the real data so a
    # whole-column merge cannot explode.
    for rng in ws.merged_cells.ranges:
        anchor = texts.get((rng.min_row, rng.min_col))
        if not anchor:
            continue
        for r in range(rng.min_row, min(rng.max_row, max_row) + 1):
            for c in range(rng.min_col, min(rng.max_col, max_col) + 1):
                texts[r, c] = anchor

    rows: dict[int, list[tuple[int, str]]] = {}
    for (r, c) in sorted(texts):
        if r in hidden_rows or c in hidden_cols:
            continue
        rows.setdefault(r, []).append((c, texts[r, c]))

    sheet = SheetLines(name=ws.title, title="")
    skipped = (len(hidden_rows), len(hidden_cols))

    header_row = next((r for r, cells in rows.items() if len(_distinct([t for _, t in cells])) >= 3), None)
    if header_row is None:
        for r, cells in rows.items():
            sheet.lines.append((r, ": ".join(_distinct([t for _, t in cells])), False))
        return sheet, *skipped

    two_row = any(
        rng.min_row == header_row and rng.max_row == header_row + 1 for rng in ws.merged_cells.ranges
    )
    top = dict(rows.get(header_row, []))
    sub = dict(rows.get(header_row + 1, [])) if two_row else {}
    names: dict[int, str] = {}
    for c in set(top) | set(sub):
        t, s = top.get(c, ""), sub.get(c, "")
        name = f"{t} - {s}" if t and s and t != s else (t or s)
        names[c] = _clean(name, MAX_NAME_CHARS, ellipsis_extra=False)
    last_header_row = header_row + 1 if two_row else header_row

    title_parts: list[str] = []
    for r, cells in rows.items():
        values = _distinct([t for _, t in cells])
        if r < header_row:
            if len(values) == 1:
                title_parts.append(values[0])
            else:
                sheet.lines.append((r, ": ".join(values), False))
        elif r > last_header_row:
            if len(values) == 1:
                sheet.lines.append((r, values[0], True))
            else:
                pairs = [f"{names.get(c) or 'Cột ' + get_column_letter(c)}: {t}" for c, t in cells]
                sheet.lines.append((r, PAIR_SEP.join(pairs), False))
    sheet.title = _clean(" — ".join(title_parts), MAX_TITLE_CHARS, ellipsis_extra=False)
    return sheet, *skipped


def read_sheets(source) -> list[SheetLines]:
    """Read the visible sheets of a .xlsx (path or binary file object) into lines.

    Raises ValueError with a Vietnamese message (stored in documents.error_message) when the
    file is not a readable .xlsx, too large to load safely, or has no readable data.
    """
    import openpyxl

    _check_zip(source)
    try:
        wb = openpyxl.load_workbook(source, data_only=True)
    except Exception:  # KeyError for a renamed .docx, InvalidFileException, XML errors, ...
        raise ValueError(INVALID_MSG) from None

    sheets: list[SheetLines] = []
    hidden_sheets = hidden_rows = hidden_cols = 0
    try:
        for ws in wb.worksheets:
            if ws.sheet_state != "visible":
                hidden_sheets += 1
                continue
            sheet, n_rows, n_cols = _read_sheet(ws)
            hidden_rows += n_rows
            hidden_cols += n_cols
            if sheet.lines:
                sheets.append(sheet)
    finally:
        wb.close()

    if hidden_sheets or hidden_rows or hidden_cols:
        logger.info(
            f"Spreadsheet: skipped {hidden_sheets} hidden sheets, {hidden_rows} hidden rows, "
            f"{hidden_cols} hidden columns"
        )
    if not sheets:
        raise ValueError(EMPTY_MSG)
    return sheets


# ---------------------------------------------------------------------------
# Chunking
# ---------------------------------------------------------------------------


def _section(name: str, first_row: int, last_row: int) -> str:
    name = _WS.sub(" ", name.replace("]", ")").replace("[", "(")).strip()[:_SHEET_NAME_MAX]
    rows = f"{first_row}" if first_row == last_row else f"{first_row}–{last_row}"
    return f"Sheet {name} › dòng {rows}"


def _split_line(text: str, budget: int) -> list[str]:
    """Split an over-long row at PAIR_SEP; continuation pieces repeat the row's first pair."""
    if len(text) <= budget:
        return [text]
    pairs = text.split(PAIR_SEP)
    lead = pairs[0] + PAIR_SEP
    if len(lead) > budget // 2:
        lead = pairs[0][: budget // 2 - len(PAIR_SEP) - 1] + "…" + PAIR_SEP

    pieces: list[str] = []
    current, has_pairs = pairs[0], False
    for pair in pairs[1:]:
        # A piece always carries at least one pair besides the lead; a lone lead says nothing.
        if not has_pairs or len(current) + len(PAIR_SEP) + len(pair) <= budget:
            current += PAIR_SEP + pair
            has_pairs = True
            continue
        pieces.append(current)
        current = lead + pair
    pieces.append(current)

    out: list[str] = []
    for piece in pieces:
        while len(piece) > budget:  # one pair alone is too long: cut it hard
            out.append(piece[:budget])
            piece = lead + piece[budget:]
        out.append(piece)
    return out


def chunk_sheets(sheets: list[SheetLines], chunk_size: int = CHUNK_SIZE) -> list[ChunkResult]:
    """Pack whole lines into chunks of about ``chunk_size`` chars, one sheet per chunk."""
    chunks: list[ChunkResult] = []

    def emit(sheet: SheetLines, head: str, lines: list[tuple[int, str, bool]]) -> None:
        section = _section(sheet.name, lines[0][0], lines[-1][0])
        body = "\n".join(text for _, text, _ in lines)
        content = f"[{section}] {head}\n{body}" if head else f"[{section}] {body}"
        chunks.append(ChunkResult(chunk_index=len(chunks), content=content, section=section))

    for sheet in sheets:
        # Fixed allowance: the longest section this sheet can produce, and the longest head.
        prefix = len(_section(sheet.name, 1048576, 1048576)) + 3
        has_label = any(is_label for _, _, is_label in sheet.lines)
        head_max = len(sheet.title) + (len(" › ") + _HEAD_LABEL_MAX + 1 if has_label else 0) + 1
        budget = max(chunk_size - prefix - head_max, _MIN_BODY)

        current: list[tuple[int, str, bool]] = []
        size = 0
        head = sheet.title
        last_label = ""
        for row, text, is_label in sheet.lines:
            for piece in _split_line(text, budget):
                if current and size + 1 + len(piece) > budget:
                    emit(sheet, head, current)
                    current, size = [], 0
                if not current:
                    carried = "" if is_label or not last_label else _clean(last_label, _HEAD_LABEL_MAX, False)
                    head = " › ".join(p for p in (sheet.title, carried) if p)
                size += len(piece) + (1 if current else 0)
                current.append((row, piece, is_label))
            if is_label:
                last_label = text
        if current:
            emit(sheet, head, current)
    if len(chunks) > MAX_CHUNKS:
        raise ValueError(_too_many_chunks_msg(len(chunks)))
    return chunks
