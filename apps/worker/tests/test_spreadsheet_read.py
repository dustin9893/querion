"""Kiểm thử `read_sheets` / `_cell_text` — đọc bảng tính .xlsx thành các dòng tự mô tả.

Sự cố mà các test này canh giữ:
- Bảng tính đi qua bộ cắt văn bản xuôi (`chunk_text`) thì mất tên cột, dòng bị cắt đôi, dòng
  đánh số ngắn bị bỏ — mỗi đoạn được embed riêng nên một dãy giá trị không tên cột là vô nghĩa.
- Ô gộp (tiêu đề, cột "Nhóm" gộp dọc, dòng tên cột hai tầng) làm lệch tên cột hoặc mất giá trị.
- Sheet/dòng/cột ẩn (cột lãi biên nội bộ) lọt vào trợ lý cho khách hàng.
- Dấu "<" ">" bị `sanitize_chunk` của guard xoá cả cụm "<...>", mất dữ liệu phí.
- Tệp quá lớn làm worker bị OOM-kill, văn bản kẹt mãi ở trạng thái "indexing"; tệp .xls cũ /
  .docx đổi đuôi hiện thông báo lỗi thô của thư viện thay vì lời nhắn tiếng Việt.
- Handle tệp bị giữ lại trên Windows (thiếu wb.close()).

Run: apps/api/.venv/Scripts/python.exe -m pytest apps/worker/tests -q
(từ gốc repo)
"""

import datetime as dt
import io
import os
import re
import zipfile

import pytest
from docx import Document
from openpyxl import Workbook

from tests.xlsx_fixtures import SHEET1, assert_no_loss, emitted_text, patch_xml, to_bytes, to_path
from worker.pipeline import spreadsheet
from worker.pipeline.spreadsheet import EMPTY_MSG, INVALID_MSG, MAX_CELL_CHARS, _cell_text, read_sheets


def _unzip(data: bytes, member: str) -> str:
    return zipfile.ZipFile(io.BytesIO(data)).read(member).decode("utf-8")


def _read(wb: Workbook):
    return read_sheets(io.BytesIO(to_bytes(wb)))


def _only(sheets):
    assert len(sheets) == 1
    return sheets[0]


# ---------------------------------------------------------------------------
# Structure: title, header, merged cells
# ---------------------------------------------------------------------------


def test_merged_title_above_header_becomes_title():
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "BIỂU PHÍ THẺ"
    ws.merge_cells("A1:D1")
    ws.append(["Mã", "Dịch vụ", "Mức phí", "Ghi chú"])
    ws.append(["P01", "Phát hành", "50000", "Miễn năm đầu"])

    sheet = _only(_read(wb))

    assert sheet.title == "BIỂU PHÍ THẺ"
    assert sheet.lines[0] == (3, "Mã: P01 | Dịch vụ: Phát hành | Mức phí: 50000 | Ghi chú: Miễn năm đầu", False)
    assert_no_loss(wb, [sheet])


def test_vertically_merged_group_is_filled_down():
    wb = Workbook()
    ws = wb.active
    ws.append(["Tiêu đề"])
    ws.append(["Nhóm", "Dịch vụ", "Mức phí"])
    ws.append(["Chuyển tiền", "Trong nước", "11000"])
    ws.append([None, "Liên ngân hàng", "22000"])
    ws.append([None, "Quốc tế", "33000"])
    ws.merge_cells("A3:A5")

    sheet = _only(_read(wb))

    assert [row for row, _, _ in sheet.lines] == [3, 4, 5]
    assert all(text.startswith("Nhóm: Chuyển tiền | ") for _, text, _ in sheet.lines)
    assert_no_loss(wb, [sheet])


def test_two_row_merged_header_combines_top_and_sub_names():
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "LÃI SUẤT TIỀN GỬI"
    ws["A3"], ws["B3"], ws["C3"] = "Kỳ hạn", "Loại", "Lãi suất"
    ws["C4"], ws["D4"] = "Cá nhân", "Doanh nghiệp"
    ws.merge_cells("A3:A4")
    ws.merge_cells("B3:B4")
    ws.merge_cells("C3:D3")
    ws.append(["6 tháng", "Thường", "4,5%", "4,2%"])

    sheet = _only(_read(wb))

    assert sheet.lines == [
        (5, "Kỳ hạn: 6 tháng | Loại: Thường | Lãi suất - Cá nhân: 4,5% | Lãi suất - Doanh nghiệp: 4,2%", False)
    ]
    assert_no_loss(wb, [sheet])


def test_blank_header_cell_is_named_by_column_letter():
    wb = Workbook()
    ws = wb.active
    ws.append(["Mã", "Tên", None, "Hạn mức"])
    ws.append(["A1", "Thẻ vàng", "ghi chú lẻ", "50 triệu"])

    sheet = _only(_read(wb))

    assert sheet.lines[0][1] == "Mã: A1 | Tên: Thẻ vàng | Cột C: ghi chú lẻ | Hạn mức: 50 triệu"


def test_two_column_sheet_without_table_keeps_key_value_pairs():
    wb = Workbook()
    ws = wb.active
    ws.append(["Phí phát hành", "50.000 VND"])
    ws.append(["Phí thường niên", "Miễn phí"])

    sheet = _only(_read(wb))

    assert sheet.title == ""
    assert sheet.lines == [(1, "Phí phát hành: 50.000 VND", False), (2, "Phí thường niên: Miễn phí", False)]
    assert_no_loss(wb, [sheet])


def test_one_column_list_emits_every_value():
    wb = Workbook()
    ws = wb.active
    for name in ["Hộ chiếu", "CCCD", "Giấy phép lái xe"]:
        ws.append([name])

    sheet = _only(_read(wb))

    assert sheet.title == ""
    assert [text for _, text, _ in sheet.lines] == ["Hộ chiếu", "CCCD", "Giấy phép lái xe"]


def test_app_exporter_layout_picks_the_table_header_not_a_summary_row():
    # Bố cục do reports.py của chính ứng dụng ghi: tiêu đề gộp, "Lập lúc …", các dòng tóm tắt
    # hai giá trị, một dòng trống, rồi bảng >= 3 cột và dòng tổng.
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Báo cáo dư nợ"
    ws.merge_cells("A1:C1")
    ws["A2"] = "Lập lúc 21/09/2026 08:00"
    ws["A4"], ws["B4"] = "Tổng dư nợ", "1.200 tỷ"
    ws["A5"], ws["B5"] = "Số khách hàng", "320"
    ws["A7"], ws["B7"], ws["C7"] = "Chi nhánh", "Dư nợ", "Nợ xấu"
    ws["A8"], ws["B8"], ws["C8"] = "Hà Nội", "700 tỷ", "1,2%"
    ws["A9"], ws["B9"] = "Tổng cộng", "1.200 tỷ"

    sheet = _only(_read(wb))

    assert sheet.title == "Báo cáo dư nợ — Lập lúc 21/09/2026 08:00"
    assert sheet.lines == [
        (4, "Tổng dư nợ: 1.200 tỷ", False),
        (5, "Số khách hàng: 320", False),
        (8, "Chi nhánh: Hà Nội | Dư nợ: 700 tỷ | Nợ xấu: 1,2%", False),
        (9, "Chi nhánh: Tổng cộng | Dư nợ: 1.200 tỷ", False),
    ]
    assert_no_loss(wb, [sheet])


def test_section_band_after_header_is_a_label_line():
    wb = Workbook()
    ws = wb.active
    ws.append(["Mã", "Dịch vụ", "Mức phí"])
    ws["A2"] = "II. CHUYỂN TIỀN QUỐC TẾ"
    ws.merge_cells("A2:C2")
    ws.append(["Q01", "SWIFT", "0,2%"])

    sheet = _only(_read(wb))

    assert sheet.lines == [(2, "II. CHUYỂN TIỀN QUỐC TẾ", True), (3, "Mã: Q01 | Dịch vụ: SWIFT | Mức phí: 0,2%", False)]


def test_empty_rows_keep_true_excel_row_numbers_and_empty_sheets_are_dropped():
    wb = Workbook()
    ws = wb.active
    ws.title = "Dữ liệu"
    ws["A2"], ws["B2"], ws["C2"] = "Mã", "Tên", "Giá"
    ws["A5"], ws["B5"], ws["C5"] = "X1", "Một", "10"
    ws["A9"], ws["B9"], ws["C9"] = "X2", "Hai", "20"
    wb.create_sheet("Trống")

    sheets = _read(wb)

    assert [s.name for s in sheets] == ["Dữ liệu"]
    assert [row for row, _, _ in sheets[0].lines] == [5, 9]


# ---------------------------------------------------------------------------
# Cell text
# ---------------------------------------------------------------------------


def _cell(value, number_format="General"):
    ws = Workbook().active
    cell = ws["A1"]
    cell.value = value
    cell.number_format = number_format
    return cell


@pytest.mark.parametrize(
    ("value", "number_format", "expected"),
    [
        (dt.date(2026, 3, 1), "General", "01/03/2026"),
        (dt.datetime(2026, 3, 1, 0, 0), "General", "01/03/2026"),
        (dt.datetime(2026, 3, 1, 14, 30), "General", "01/03/2026 14:30"),
        (dt.time(8, 5), "General", "08:05"),
        (0.0002, "0.00%", "0,02%"),
        (0.952, "0.0%", "95,2%"),
        (1, "0%", "100%"),
        (1234567, "#,##0", "1.234.567"),
        (1234567.5, "#,##0.00", "1.234.567,5"),
        (1234567, '#,##0 "VND"', "1.234.567"),
        (1234567, "[$-42A]#,##0;[Red]-#,##0", "1.234.567"),
        (-2500000, "#,##0", "-2.500.000"),
        (1234567, "General", "1234567"),
        (1001, "General", "1001"),
        (2026, "0", "2026"),
        (0.1 + 0.2, "General", "0,3"),
        (5.0, "General", "5"),
        (-1.25, "General", "-1,25"),
        (True, "General", "Có"),
        (False, "General", "Không"),
        ("#DIV/0!", "General", ""),
        (None, "General", ""),
    ],
)
def test_cell_text(value, number_format, expected):
    assert _cell_text(_cell(value, number_format)) == expected


def test_error_cell_is_an_error_type():
    # Bảo đảm dòng "#DIV/0!" ở trên thật sự là ô lỗi, không phải chuỗi thường.
    assert _cell("#DIV/0!").data_type == "e"


def test_angle_brackets_become_fullwidth_signs():
    wb = Workbook()
    ws = wb.active
    ws.append(["Hạn mức", "Phí", "Tối thiểu"])
    ws.append(["<5 triệu", "0,1%", ">10 USD"])

    text = emitted_text(_read(wb))

    assert "＜5 triệu" in text and "＞10 USD" in text
    assert "<" not in text and ">" not in text


def test_long_cell_is_capped_with_ellipsis():
    text = _cell_text(_cell("x" * 600))

    assert len(text) <= MAX_CELL_CHARS + 1
    assert text.endswith("…")


def test_multiline_cell_stays_on_one_line():
    wb = Workbook()
    ws = wb.active
    ws.append(["Mã", "Điều kiện", "Phí"])
    ws.append(["P1", "Dòng một\nDòng hai\r\n  Dòng ba", "10"])

    sheet = _only(_read(wb))

    assert len(sheet.lines) == 1
    assert "\n" not in sheet.lines[0][1]
    assert "Điều kiện: Dòng một Dòng hai Dòng ba" in sheet.lines[0][1]


# ---------------------------------------------------------------------------
# Formulas
# ---------------------------------------------------------------------------


def _formula_workbook() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(["Mã", "Số lượng", "Tổng"])
    ws.append(["F1", "3", "=SUM(B2:B2)"])
    return to_bytes(wb)


def test_formula_with_cached_value_is_indexed():
    data = patch_xml(_formula_workbook(), SHEET1, lambda xml: xml.replace("<v></v>", "<v>12345</v>", 1))
    assert "<v>12345</v>" in _unzip(data, SHEET1)

    text = emitted_text(read_sheets(io.BytesIO(data)))

    assert "Tổng: 12345" in text


def test_formula_without_cached_value_is_skipped_and_row_kept():
    sheet = _only(read_sheets(io.BytesIO(_formula_workbook())))

    assert sheet.lines == [(2, "Mã: F1 | Số lượng: 3", False)]
    assert "SUM" not in emitted_text([sheet])


# ---------------------------------------------------------------------------
# Hidden content never reaches the index
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("state", ["hidden", "veryHidden"])
def test_hidden_sheets_are_skipped(state):
    wb = Workbook()
    wb.active.append(["Mã", "Tên", "Giá"])
    wb.active.append(["V1", "Hiện", "10"])
    secret = wb.create_sheet("Nội bộ")
    secret.append(["LAI_BIEN_SECRET", "x", "y"])
    secret.sheet_state = state

    text = emitted_text(_read(wb))

    assert "LAI_BIEN_SECRET" not in text
    assert "V1" in text


def test_hidden_row_is_skipped():
    wb = Workbook()
    ws = wb.active
    ws.append(["Mã", "Tên", "Giá"])
    ws.append(["V1", "Hiện", "10"])
    ws.append(["H1", "HIDDEN_ROW_SECRET", "20"])
    ws.append(["V2", "Hiện nữa", "30"])
    ws.row_dimensions[3].hidden = True

    sheet = _only(_read(wb))

    assert "HIDDEN_ROW_SECRET" not in emitted_text([sheet])
    assert [row for row, _, _ in sheet.lines] == [2, 4]


def test_hidden_single_column_is_skipped():
    wb = Workbook()
    ws = wb.active
    ws.append(["Mã", "Tên", "Lãi biên", "Giá"])
    ws.append(["V1", "Hiện", "MARGIN_SECRET", "10"])
    ws.column_dimensions["C"].hidden = True

    text = emitted_text(_read(wb))

    assert "MARGIN_SECRET" not in text and "Lãi biên" not in text
    assert "Mã: V1 | Tên: Hiện | Giá: 10" in text


def test_hidden_column_group_is_skipped():
    wb = Workbook()
    ws = wb.active
    ws.append(["Mã", "Tên", "Ẩn 1", "Ẩn 2", "Ẩn 3", "Giá"])
    ws.append(["V1", "Hiện", "SECRET_C", "SECRET_D", "SECRET_E", "10"])
    # Excel lưu nhóm cột ẩn C:E một lần: <col min="3" max="5" hidden="1"/>.
    data = patch_xml(
        to_bytes(wb), SHEET1, lambda xml: xml.replace("<sheetData>", '<cols><col min="3" max="5" hidden="1"/></cols><sheetData>', 1)
    )
    assert '<col min="3" max="5" hidden="1"/>' in _unzip(data, SHEET1)

    text = emitted_text(read_sheets(io.BytesIO(data)))

    assert not re.search(r"SECRET|Ẩn", text)
    assert "Mã: V1 | Tên: Hiện | Giá: 10" in text


# ---------------------------------------------------------------------------
# Errors: Vietnamese messages, never raw library text
# ---------------------------------------------------------------------------


def _simple_workbook() -> Workbook:
    wb = Workbook()
    wb.active.append(["Mã", "Tên", "Giá"])
    wb.active.append(["V1", "Một", "10"])
    return wb


def test_too_large_xml_fails_before_loading(monkeypatch):
    monkeypatch.setattr(spreadsheet, "MAX_XML_BYTES", 1024)

    def boom(*args, **kwargs):
        raise AssertionError("load_workbook must not run")

    monkeypatch.setattr("openpyxl.load_workbook", boom)

    with pytest.raises(ValueError, match=r"^Bảng tính quá lớn để lập chỉ mục .*Hãy tách thành nhiều tệp nhỏ hơn\.$"):
        _read(_simple_workbook())


def test_empty_workbook_fails_with_empty_message():
    with pytest.raises(ValueError) as exc:
        _read(Workbook())

    assert str(exc.value) == EMPTY_MSG


def test_only_hidden_data_fails_with_empty_message():
    wb = _simple_workbook()
    wb.active.row_dimensions[1].hidden = True
    wb.active.row_dimensions[2].hidden = True

    with pytest.raises(ValueError) as exc:
        _read(wb)

    assert str(exc.value) == EMPTY_MSG


def _docx_bytes() -> bytes:
    out = io.BytesIO()
    doc = Document()
    doc.add_paragraph("Quy trình cho vay")
    doc.save(out)
    return out.getvalue()


@pytest.mark.parametrize(
    "data",
    [
        pytest.param(_docx_bytes(), id="docx-renamed"),
        pytest.param(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 512, id="ole2-xls"),
        pytest.param(b"", id="zero-bytes"),
        pytest.param(b"PK\x03\x04 truncated", id="broken-zip"),
    ],
)
def test_invalid_file_fails_with_invalid_message(tmp_path, data):
    path = tmp_path / "bang.xlsx"
    path.write_bytes(data)

    with pytest.raises(ValueError) as exc:
        read_sheets(str(path))

    assert str(exc.value) == INVALID_MSG


def test_file_handle_is_released_after_reading(tmp_path):
    path = to_path(_simple_workbook(), tmp_path)

    read_sheets(path)
    os.remove(path)

    assert not os.path.exists(path)


def test_file_handle_is_released_after_invalid_file(tmp_path):
    path = tmp_path / "bang.xlsx"
    path.write_bytes(_docx_bytes())

    with pytest.raises(ValueError):
        read_sheets(str(path))
    os.remove(path)

    assert not path.exists()
