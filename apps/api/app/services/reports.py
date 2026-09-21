"""Turning a workflow run into a file: Markdown or a filled .docx template.

Templates are written by admins of a unit (not by the model and not by staff), but they are
still rendered in a **sandboxed** Jinja environment: a template is data in the database, and a
compromised admin account must not be able to read files or reach Python internals through it.

The rendered bytes go to MinIO and one ``artifacts`` row records what was produced, for which
run, and who may download it.
"""

from __future__ import annotations

import logging
import re
import uuid
from datetime import date, datetime, timedelta, timezone
from typing import Any

from jinja2 import StrictUndefined
from jinja2.sandbox import SandboxedEnvironment
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.artifact import Artifact
from app.storage import upload_file

logger = logging.getLogger(__name__)

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MARKDOWN_MIME = "text/markdown; charset=utf-8"
MAX_PREVIEW_CHARS = 20000
DEFAULT_RETENTION_DAYS = 30
VN_TZ = timezone(timedelta(hours=7))


class ReportError(Exception):
    """The template could not be rendered. The message is shown to the admin, not to customers."""


# ---------------------------------------------------------------------------
# Template rendering
# ---------------------------------------------------------------------------

def _fmt_number(value: Any, decimals: int = 0) -> str:
    """1234567.5 -> '1.234.568' (Vietnamese grouping: dot for thousands, comma for decimals)."""
    try:
        number = float(value)
    except (TypeError, ValueError):
        return "" if value is None else str(value)
    whole, _, frac = f"{number:,.{decimals}f}".partition(".")
    whole = whole.replace(",", ".")
    return f"{whole},{frac}" if frac else whole


def _fmt_date(value: Any, pattern: str = "%d/%m/%Y") -> str:
    if isinstance(value, datetime):
        return value.astimezone(VN_TZ).strftime(pattern)
    if isinstance(value, str):
        for source in ("%Y-%m-%d", "%Y-%m-%dT%H:%M:%S", "%d/%m/%Y"):
            try:
                return datetime.strptime(value[:len(datetime.now().strftime(source))], source).strftime(pattern)
            except ValueError:
                continue
    return "" if value is None else str(value)


def jinja_env() -> SandboxedEnvironment:
    env = SandboxedEnvironment(autoescape=False, undefined=StrictUndefined)
    env.filters["so"] = _fmt_number          # number with Vietnamese grouping
    env.filters["tien"] = lambda v: _fmt_number(v) + " đ"
    env.filters["ngay"] = _fmt_date
    return env


def build_context(state: dict[str, Any], *, extra: dict[str, Any] | None = None) -> dict[str, Any]:
    """Everything a report template may read. Keep the names stable: templates are user data."""
    now = datetime.now(VN_TZ)
    ctx: dict[str, Any] = {
        "query": state.get("query", ""),
        "inputs": state.get("inputs", {}) or {},
        "params": state.get("extracted_params", {}) or {},
        "tool_results": state.get("tool_results", {}) or {},
        "answer": state.get("answer", "") or "",
        "chunks": state.get("retrieved_chunks", []) or [],
        "now": now,
        "today": now.strftime("%d/%m/%Y"),
        "time": now.strftime("%H:%M"),
    }
    # Inputs are also exposed at the top level so a template can just write {{ tu_ngay }}.
    for key, value in (state.get("inputs") or {}).items():
        ctx.setdefault(str(key), value)
    ctx.update(extra or {})
    return ctx


def render_markdown(template: str, context: dict[str, Any]) -> str:
    try:
        return jinja_env().from_string(template or "").render(**context)
    except Exception as exc:
        raise ReportError(f"Lỗi mẫu Markdown: {type(exc).__name__}: {exc}")


def render_docx(template_bytes: bytes, context: dict[str, Any]) -> bytes:
    from io import BytesIO

    try:
        from docxtpl import DocxTemplate
    except ImportError:  # pragma: no cover
        raise ReportError("Thiếu thư viện docxtpl trên máy chủ")
    try:
        doc = DocxTemplate(BytesIO(template_bytes))
        doc.render(context, jinja_env())
        out = BytesIO()
        doc.save(out)
        return out.getvalue()
    except ReportError:
        raise
    except Exception as exc:
        raise ReportError(f"Lỗi mẫu DOCX: {type(exc).__name__}: {exc}")


def docx_variables(template_bytes: bytes) -> set[str]:
    """Placeholders a .docx template expects — used to check a template on upload."""
    from io import BytesIO

    from docxtpl import DocxTemplate

    try:
        return set(DocxTemplate(BytesIO(template_bytes)).get_undeclared_template_variables(jinja_env()))
    except Exception as exc:
        raise ReportError(f"Không đọc được mẫu DOCX: {type(exc).__name__}: {exc}")


# ---------------------------------------------------------------------------
# Spreadsheets
# ---------------------------------------------------------------------------

# Excel number formats per column "format" (VND has no decimals in practice).
CELL_FORMATS = {
    "money": '#,##0" đ"',
    "number": "#,##0.##",
    "int": "#,##0",
    "percent": "0.0%",
}
MAX_SHEET_ROWS = 20000


def _cell_value(row: Any, column: dict) -> Any:
    """Read one cell out of a data row: `field` is a dot path, or `value` is a fixed string."""
    if "value" in column:
        return column["value"]
    path = str(column.get("field") or "")
    current: Any = row
    for part in path.split("."):
        if part == "":
            continue
        if isinstance(current, dict):
            current = current.get(part)
        elif isinstance(current, list) and part.isdigit() and int(part) < len(current):
            current = current[int(part)]
        else:
            return None
    return current


# openpyxl refuses these in a sheet title, and a model naming a sheet "Doanh số 08/2026" would
# otherwise crash the whole export instead of producing a file.
_BAD_SHEET_CHARS = re.compile(r"[\\/*?:\[\]]")


def safe_sheet_name(name: Any, index: int, used: set[str]) -> str:
    """A title Excel accepts: no \ / * ? : [ ], at most 31 chars, never empty, never repeated."""
    title = _BAD_SHEET_CHARS.sub("-", str(name or "")).strip().strip("'")[:31].strip()
    if not title:
        title = f"Sheet{index + 1}"
    base, n = title, 2
    while title.casefold() in used:
        suffix = f" ({n})"
        title, n = f"{base[:31 - len(suffix)]}{suffix}", n + 1
    used.add(title.casefold())
    return title


def _put(ws, row: int, column: int, value: Any):
    """Write one cell, never a formula.

    Cell content can come from a retrieved document, a tool result or the model itself, and
    openpyxl turns any string starting with "=" into a formula. Forcing the type back to string
    keeps the text visible and stops a spreadsheet from executing what a document said. A value
    openpyxl cannot store (a dict, a list) becomes its text instead of killing the whole file.
    """
    if not isinstance(value, (str, int, float, bool, datetime, date, type(None))):
        value = str(value)
    cell = ws.cell(row=row, column=column, value=value)
    if cell.data_type == "f":
        cell.data_type = "s"
    return cell



# ---------------------------------------------------------------------------
# Charts
# ---------------------------------------------------------------------------

CHART_TYPES = {"bar", "bar_h", "line", "pie"}
MAX_CHARTS_PER_SHEET = 3
MAX_CHART_SERIES = 4


def _chart_specs(sheet: dict, columns: list[dict]) -> list[dict]:
    """Normalise ``sheet["charts"]`` into ``{type, title, category, series[], y_format}``.

    A chart names *columns* of the same sheet (by ``field``), never raw cell ranges: the data is
    already laid out by ``render_xlsx``, so the chart can only ever point at cells that exist.
    Unknown fields and unknown types are dropped silently — a report with a slightly wrong chart
    spec must still produce its table.
    """
    fields = {str(c.get("field") or c.get("header") or i): i for i, c in enumerate(columns)}
    out: list[dict] = []
    for raw in (sheet.get("charts") or [])[:MAX_CHARTS_PER_SHEET]:
        if not isinstance(raw, dict):
            continue
        kind = str(raw.get("type") or "bar").lower()
        if kind not in CHART_TYPES:
            continue
        category = fields.get(str(raw.get("category_field") or ""))
        series: list[tuple[int, str | None]] = []
        for item in (raw.get("series") or [])[:MAX_CHART_SERIES]:
            field, name = (item.get("field"), item.get("name")) if isinstance(item, dict) else (item, None)
            index = fields.get(str(field or ""))
            if index is not None and index != category:
                series.append((index, str(name) if name else None))
        if category is None or not series:
            continue
        if kind == "pie":
            series = series[:1]
        y_format = CELL_FORMATS.get(str(columns[series[0][0]].get("format") or ""), "")
        out.append({"type": kind, "title": str(raw.get("title") or "")[:120], "category": category,
                    "series": series, "y_format": y_format, "y_title": str(raw.get("y_title") or "")[:60]})
    return out


def _add_charts(ws, specs: list[dict], *, header_line: int, first_row: int, last_row: int,
                anchor_col: int) -> None:
    """Draw the charts of one sheet to the right of its table, stacked top to bottom."""
    if last_row < first_row or not specs:
        return
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference
    from openpyxl.chart.label import DataLabelList
    from openpyxl.utils import get_column_letter

    row = header_line
    for spec in specs:
        if spec["type"] == "pie":
            chart = PieChart()
            chart.dataLabels = DataLabelList()
            chart.dataLabels.showPercent = True
        elif spec["type"] == "line":
            chart = LineChart()
        else:
            chart = BarChart()
            chart.type = "bar" if spec["type"] == "bar_h" else "col"
            chart.grouping = "clustered"
        chart.title = spec["title"] or None
        chart.height, chart.width = 7.5, 16
        chart.style = 10
        for col_index, name in spec["series"]:
            # the header cell gives the series its name, so titles_from_data reads one extra row
            ref = Reference(ws, min_col=col_index + 1, min_row=header_line, max_row=last_row)
            chart.add_data(ref, titles_from_data=True)
            if name and chart.series:
                from openpyxl.chart.series import SeriesLabel
                chart.series[-1].tx = SeriesLabel(v=name)
        chart.set_categories(Reference(ws, min_col=spec["category"] + 1, min_row=first_row, max_row=last_row))
        if spec["type"] != "pie":
            if spec["y_format"]:
                chart.y_axis.number_format = spec["y_format"]
            if spec["y_title"]:
                chart.y_axis.title = spec["y_title"]
            chart.y_axis.majorGridlines = chart.y_axis.majorGridlines  # keep defaults explicit
            chart.legend.position = "b"
        ws.add_chart(chart, f"{get_column_letter(anchor_col)}{row}")
        row += 16


def render_xlsx(sheets: list[dict], context: dict[str, Any], *, title: str | None = None) -> bytes:
    """Build a .xlsx from a declarative spec — a report is a table, not a text template.

    Each sheet: ``{name, title?, rows: [...], columns: [{header, field|value, format?, width?,
    total?}], summary?: [[label, value], ...], charts?: [{type: bar|bar_h|line|pie, title,
    category_field, series: [{field, name?}], y_title?}]}``. `rows` is already resolved to a real
    list by the node (a lone ``{{...}}`` placeholder keeps its type), so nothing is stringified on
    the way in. Charts reference columns of the same sheet, never cell ranges (see `_chart_specs`).
    """
    from io import BytesIO

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:  # pragma: no cover
        raise ReportError("Thiếu thư viện openpyxl trên máy chủ")

    if not sheets:
        raise ReportError("Node chưa khai báo sheet nào cho tệp Excel")

    wb = Workbook()
    wb.remove(wb.active)
    used_titles: set[str] = set()
    header_fill = PatternFill("solid", fgColor="EE6D1F")
    header_font = Font(bold=True, color="FFFFFF")

    for index, sheet in enumerate(sheets):
        if not isinstance(sheet, dict):
            raise ReportError("Mỗi sheet phải là object")
        columns = [c for c in (sheet.get("columns") or []) if isinstance(c, dict)]
        if not columns:
            raise ReportError(f"Sheet '{sheet.get('name') or index + 1}' chưa khai báo cột")
        rows = sheet.get("rows")
        if isinstance(rows, dict):
            rows = [rows]
        rows = [r for r in (rows or []) if r is not None][:MAX_SHEET_ROWS]

        ws = wb.create_sheet(safe_sheet_name(sheet.get("name"), index, used_titles))
        line = 1
        sheet_title = sheet.get("title") or (title if index == 0 else None)
        if sheet_title:
            _put(ws, 1, 1, str(sheet_title)).font = Font(bold=True, size=14)
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(2, len(columns)))
            ws.cell(row=2, column=1, value=f"Lập lúc {context.get('today', '')} {context.get('time', '')}").font = \
                Font(italic=True, size=9)
            line = 4

        for summary_row in (sheet.get("summary") or []):
            if isinstance(summary_row, (list, tuple)) and summary_row:
                _put(ws, line, 1, str(summary_row[0])).font = Font(bold=True)
                if len(summary_row) > 1:
                    _put(ws, line, 2, summary_row[1])
                line += 1
        if sheet.get("summary"):
            line += 1

        header_line = line
        for col_index, column in enumerate(columns, start=1):
            cell = _put(ws, header_line, col_index, str(column.get("header") or column.get("field") or ""))
            cell.fill, cell.font = header_fill, header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        line += 1

        for row in rows:
            for col_index, column in enumerate(columns, start=1):
                value = _cell_value(row, column)
                cell = _put(ws, line, col_index, value)
                fmt = CELL_FORMATS.get(str(column.get("format") or ""))
                if fmt and isinstance(value, (int, float)) and not isinstance(value, bool):
                    cell.number_format = fmt
            line += 1

        if any(c.get("total") for c in columns) and rows:
            for col_index, column in enumerate(columns, start=1):
                cell = ws.cell(row=line, column=col_index)
                if col_index == 1:
                    cell.value = "Tổng cộng"
                elif column.get("total"):
                    values = [_cell_value(r, column) for r in rows]
                    cell.value = sum(v for v in values if isinstance(v, (int, float)) and not isinstance(v, bool))
                    fmt = CELL_FORMATS.get(str(column.get("format") or ""))
                    if fmt:
                        cell.number_format = fmt
                cell.font = Font(bold=True)

        # charts sit to the right of the table so the data stays the first thing a reader sees
        _add_charts(ws, _chart_specs(sheet, columns), header_line=header_line,
                    first_row=header_line + 1, last_row=header_line + len(rows), anchor_col=len(columns) + 2)

        ws.freeze_panes = ws.cell(row=header_line + 1, column=1)
        for col_index, column in enumerate(columns, start=1):
            width = column.get("width")
            if not width:
                longest = max([len(str(column.get("header") or ""))]
                              + [len(str(_cell_value(r, column) or "")) for r in rows[:200]] or [10])
                width = min(48, max(10, longest + 2))
            ws.column_dimensions[get_column_letter(col_index)].width = width

    out = BytesIO()
    wb.save(out)
    return out.getvalue()


# ---------------------------------------------------------------------------
# Storing the result
# ---------------------------------------------------------------------------

_SAFE_NAME = re.compile(r"[^0-9A-Za-zÀ-ỹ._ -]+")


def safe_filename(name: str, extension: str) -> str:
    base = _SAFE_NAME.sub("", (name or "bao-cao").strip())[:120].strip() or "bao-cao"
    if not base.lower().endswith(f".{extension}"):
        base = f"{base}.{extension}"
    return base


def template_storage_key(workspace_id, workflow_id, filename: str) -> str:
    return f"templates/{workspace_id}/{workflow_id}/{uuid.uuid4().hex}-{safe_filename(filename, 'docx')}"


async def store_artifact(
    db: AsyncSession,
    *,
    workspace_id,
    data: bytes,
    filename: str,
    content_type: str,
    title: str | None = None,
    kind: str = "report",
    audience: str = "admin",
    run_id=None,
    workflow_id=None,
    schedule_id=None,
    preview: str | None = None,
    created_by_user_id=None,
    created_by_employee_id=None,
    retention_days: int | None = DEFAULT_RETENTION_DAYS,
) -> Artifact:
    """Upload the bytes to MinIO and record one artifact row (not committed)."""
    artifact_id = uuid.uuid4()
    key = f"reports/{workspace_id}/{artifact_id}/{filename}"
    upload_file(key, data, content_type)
    artifact = Artifact(
        id=artifact_id, workspace_id=workspace_id, run_id=run_id, workflow_id=workflow_id,
        schedule_id=schedule_id, kind=kind, audience=audience, title=title or filename,
        filename=filename, content_type=content_type, storage_key=key, size=len(data),
        preview=(preview or "")[:MAX_PREVIEW_CHARS] or None,
        created_by_user_id=created_by_user_id, created_by_employee_id=created_by_employee_id,
        expires_at=(datetime.now(timezone.utc) + timedelta(days=retention_days)) if retention_days else None,
    )
    db.add(artifact)
    return artifact
