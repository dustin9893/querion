"""Biểu đồ trong tệp Excel và trong công cụ xuất Excel.

Một đặc tả biểu đồ sai không được làm hỏng bảng: bảng là thứ người đọc cần trước, biểu đồ là
phần thêm. Vì thế mọi nhánh lỗi ở đây phải kết thúc bằng một workbook mở được.
"""

from io import BytesIO

import jsonschema
import pytest
from openpyxl import load_workbook

from app.services.reports import build_context, render_xlsx
from app.services.tools.builtin import kiem_tra_kha_nang_tra_no
from app.services.tools.export import build_sheets, export_args_schema

ROWS = [{"cn": "CN Hà Nội", "gn": 412, "kh": 450, "tl": 0.91},
        {"cn": "CN Hồ Chí Minh", "gn": 528, "kh": 500, "tl": 1.05}]
COLUMNS = [{"header": "Chi nhánh", "field": "cn"},
           {"header": "Giải ngân", "field": "gn", "format": "money", "total": True},
           {"header": "Kế hoạch", "field": "kh", "format": "money"},
           {"header": "Hoàn thành", "field": "tl", "format": "percent"}]


def _sheet(charts):
    data = render_xlsx([{"name": "Doanh số", "rows": ROWS, "columns": COLUMNS, "charts": charts}],
                       build_context({}), title="T")
    return load_workbook(BytesIO(data))["Doanh số"]


def test_three_chart_kinds_are_drawn():
    ws = _sheet([
        {"type": "bar", "title": "GN vs KH", "category_field": "cn",
         "series": [{"field": "gn", "name": "Giải ngân"}, {"field": "kh"}]},
        {"type": "pie", "title": "Cơ cấu", "category_field": "cn", "series": ["gn"]},
        {"type": "line", "category_field": "cn", "series": ["tl"]},
    ])
    kinds = sorted(type(c).__name__ for c in ws._charts)
    assert kinds == ["BarChart", "LineChart", "PieChart"]
    # the table is still there, with its totals row
    assert ws["A4"].value == "Chi nhánh" and ws["A7"].value == "Tổng cộng"


def test_bad_chart_specs_are_dropped_not_fatal():
    ws = _sheet([
        {"type": "bogus", "category_field": "cn", "series": ["gn"]},      # loại không có
        {"type": "bar", "category_field": "khong_co", "series": ["gn"]},  # cột nhãn không tồn tại
        {"type": "bar", "category_field": "cn", "series": ["khong_co"]},  # cột số không tồn tại
        {"type": "bar", "category_field": "cn", "series": ["cn"]},        # chuỗi trùng cột nhãn
        "không phải object",
    ])
    assert ws._charts == []
    assert ws["A5"].value == "CN Hà Nội"


def test_pie_keeps_only_first_series_and_cap_applies():
    ws = _sheet([{"type": "pie", "category_field": "cn", "series": ["gn", "kh", "tl"]}] * 5)
    assert len(ws._charts) == 3  # MAX_CHARTS_PER_SHEET
    assert all(len(c.series) == 1 for c in ws._charts)


def test_no_rows_means_no_chart():
    data = render_xlsx([{"name": "S", "rows": [], "columns": COLUMNS,
                         "charts": [{"type": "bar", "category_field": "cn", "series": ["gn"]}]}],
                       build_context({}))
    assert load_workbook(BytesIO(data))["S"]._charts == []


def test_export_tool_maps_bieu_do_to_chart():
    args = {"sheets": [{
        "ten": "Doanh số", "cot": [{"nhan": "Chi nhánh", "khoa": "cn"}, {"nhan": "GN", "khoa": "gn", "dinh_dang": "tien"}],
        "dong": [{"cn": "HN", "gn": "412.000.000"}, {"cn": "HCM", "gn": "528.000.000"}],
        "bieu_do": {"loai": "tron", "tieu_de": "Cơ cấu", "truc_nhan": "cn", "chuoi": ["gn"]},
    }]}
    jsonschema.validate(args, export_args_schema())
    sheets, *_ = build_sheets(args)
    assert sheets[0]["charts"] == [{"type": "pie", "title": "Cơ cấu", "category_field": "cn", "series": ["gn"]}]
    ws = load_workbook(BytesIO(render_xlsx(sheets, build_context({}))))["Doanh số"]
    assert [type(c).__name__ for c in ws._charts] == ["PieChart"]


def test_export_tool_rejects_unknown_chart_kind():
    args = {"sheets": [{"ten": "S", "cot": [{"nhan": "a", "khoa": "a"}], "dong": [{"a": 1}],
                        "bieu_do": {"loai": "radar", "truc_nhan": "a", "chuoi": ["a"]}}]}
    with pytest.raises(jsonschema.ValidationError):
        jsonschema.validate(args, export_args_schema())


def test_dti_builtin():
    out = kiem_tra_kha_nang_tra_no(40_000_000, 5_000_000, 1_500_000_000, 9.5, 240)
    assert out["dat_nguong"] is True and 40 < out["dti_phan_tram"] < 55
    assert out["so_tien_vay_toi_da_uoc_tinh"] > 1_500_000_000
    tight = kiem_tra_kha_nang_tra_no(15_000_000, 8_000_000, 1_500_000_000, 9.5, 240)
    assert tight["dat_nguong"] is False
    assert "loi" in kiem_tra_kha_nang_tra_no(0, 0, 1, 9.5, 12)
