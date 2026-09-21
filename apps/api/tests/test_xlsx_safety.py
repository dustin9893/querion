"""Excel is strict about sheet titles and cell types, and the spreadsheet spec is written by the
model from whatever the conversation holds. A name like "Doanh số 08/2026" used to crash the whole
export, so the file came back as an error and the assistant pasted the table as text instead.

Run: cd apps/api && pytest tests/test_xlsx_safety.py -v
"""

import io

import pytest
from openpyxl import load_workbook

from app.services.reports import render_xlsx, safe_sheet_name

CONTEXT = {"today": "18/09/2026", "time": "12:00"}


def _sheet(name, rows, columns=None):
    return {"name": name, "columns": columns or [{"header": "A", "field": "a"}], "rows": rows}


def _read(data: bytes):
    return load_workbook(io.BytesIO(data))


@pytest.mark.parametrize("raw, expected", [
    ("Doanh số 08/2026", "Doanh số 08-2026"),   # the one that broke in production
    ("A\\B", "A-B"),
    ("Nhóm [1]", "Nhóm -1-"),
    ("KPI: cán bộ", "KPI- cán bộ"),
    ("Hỏi?", "Hỏi-"),
    ("  ", "Sheet1"),
    (None, "Sheet1"),
])
def test_illegal_titles_are_rewritten(raw, expected):
    assert safe_sheet_name(raw, 0, set()) == expected


def test_a_long_title_is_cut_to_excels_limit():
    assert len(safe_sheet_name("x" * 60, 0, set())) == 31


def test_repeated_titles_get_a_suffix():
    used: set[str] = set()
    assert safe_sheet_name("Doanh số", 0, used) == "Doanh số"
    assert safe_sheet_name("Doanh số", 1, used) == "Doanh số (2)"
    assert safe_sheet_name("doanh số", 2, used) == "doanh số (3)"


def test_a_workbook_with_a_slash_in_the_name_is_produced():
    data = render_xlsx([_sheet("Doanh số 08/2026", [{"a": "CN Hà Nội"}])], CONTEXT)
    assert _read(data).sheetnames == ["Doanh số 08-2026"]


def test_two_sheets_of_the_same_name_both_survive():
    data = render_xlsx([_sheet("Kỳ 08/2026", [{"a": 1}]), _sheet("Kỳ 08/2026", [{"a": 2}])], CONTEXT)
    assert _read(data).sheetnames == ["Kỳ 08-2026", "Kỳ 08-2026 (2)"]


def test_a_cell_openpyxl_cannot_store_becomes_text():
    data = render_xlsx([_sheet("S", [{"a": {"x": 1}}, {"a": [1, 2]}])], CONTEXT)
    values = [c.value for row in _read(data)["S"].iter_rows() for c in row if c.value is not None]
    assert "{'x': 1}" in values and "[1, 2]" in values


def test_a_cell_starting_with_equals_stays_text():
    data = render_xlsx([_sheet("S", [{"a": '=HYPERLINK("http://x","click")'}])], CONTEXT)
    cells = [c for row in _read(data)["S"].iter_rows() for c in row if c.value is not None]
    assert all(c.data_type != "f" for c in cells)
