"""Công cụ xuất Excel dùng chung: mô hình dựng bảng từ nội dung hội thoại rồi trả tệp tải được.

Khác smoke_chat_reports (chạy luồng báo cáo dựng sẵn), ở đây không có luồng nào cả — bảng đến từ
đúng những gì vừa trao đổi trong chat.

Cần: API :8000, mock core :8095, provider LLM, đã chạy `python -m app.seed_demo`.
Chạy:  apps/api/.venv/bin/python scratch/smoke_export.py
       API=https://<domain> ADMIN_PASSWORD=... apps/api/.venv/bin/python scratch/smoke_export.py
"""
import io
import json
import os
import sys
import uuid

import httpx
from openpyxl import load_workbook

API = os.environ.get("API", "http://localhost:8000")
c = httpx.Client(base_url=API, timeout=300)
results = []
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
created: list[str] = []


def ok(label, cond, extra=""):
    results.append((label, bool(cond)))
    print(("PASS " if cond else "FAIL ") + label, "" if cond else "→ " + str(extra)[:300])
    return bool(cond)


def sse(resp):
    out = []
    for line in resp.iter_lines():
        if line.startswith("data: ") and line[6:] != "[DONE]":
            try:
                out.append(json.loads(line[6:]))
            except Exception:
                pass
    return out


def cells(content: bytes, sheet=None):
    wb = load_workbook(io.BytesIO(content))
    ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
    return wb, ws, [cell for row in ws.iter_rows() for cell in row if cell.value is not None]


tok = c.post("/v1/auth/login", json={"email": "admin@querion.io",
                                     "password": os.environ.get("ADMIN_PASSWORD", "admin123")}).json()["access_token"]
A = {"Authorization": f"Bearer {tok}"}
ws_ids = {w["name"]: w["id"] for w in c.get("/v1/audit/filters", headers=A).json()["workspaces"]}
EB = {**A, "X-Workspace-Id": ws_ids["Khối Khách hàng Doanh nghiệp (EB)"]}
# any other unit: the export tool is shared bank-wide, so its name does not matter
other_ws = next((i for n, i in ws_ids.items() if n != "Khối Khách hàng Doanh nghiệp (EB)"), None)
OTHER = {**A, "X-Workspace-Id": other_ws} if other_ws else None
tag = uuid.uuid4().hex[:6]

# ---- 1. công cụ có sẵn sau seed
tools = {t["slug"]: t for t in c.get("/v1/tools", headers=EB).json()}
ex = tools.get("xuat_excel")
if not ok("seed tạo công cụ xuat_excel", ex is not None, list(tools)):
    sys.exit(1)
ok("đúng loại 'export' và mở cho toàn ngân hàng", ex["kind"] == "export" and ex["share_scope"] == "bank",
   (ex["kind"], ex["share_scope"]))
ok("đơn vị khác cũng dùng được (chia sẻ toàn ngân hàng)",
   OTHER is not None and any(t["slug"] == "xuat_excel" for t in c.get("/v1/tools", headers=OTHER).json()),
   "không có đơn vị thứ hai để kiểm tra" if OTHER is None else "")
schema = ex["params_schema"]
ok("tham số do hệ thống quy định: khai báo sheets/cột/dòng",
   "sheets" in schema.get("properties", {})
   and "cot" in schema["properties"]["sheets"]["items"]["properties"]
   and "dong" in schema["properties"]["sheets"]["items"]["properties"], list(schema.get("properties", {})))
ok("không cần chọn khách hàng / duyệt", ex["allow_customer"] is False and ex["requires_approval"] is False)

# ---- 2. tạo công cụ export: định dạng lạ bị chặn, schema tự đặt bị bỏ qua
bad = c.post("/v1/tools", headers=EB, json={"slug": f"xuat_pdf_{tag}", "name": "Xuất PDF", "kind": "export",
                                            "description": "Xuất bảng ra PDF", "config": {"format": "pdf"}})
ok("định dạng ngoài xlsx bị chặn (400)", bad.status_code == 400 and "xlsx" in bad.text, bad.status_code)
mine = c.post("/v1/tools", headers=EB, json={"slug": f"xuat_excel_{tag}", "name": f"Xuất Excel {tag}",
                                             "kind": "export", "description": "Xuất bảng trong hội thoại ra Excel",
                                             "config": {"format": "xlsx"},
                                             "params_schema": {"type": "object", "properties": {"hack": {"type": "string"}}}})
ok("tạo được công cụ export mới (201)", mine.status_code == 201, mine.text[:200])
new_tool = mine.json() if mine.status_code == 201 else {}
if new_tool:
    ok("schema tự đặt bị bỏ qua, dùng schema chuẩn",
       "hack" not in new_tool["params_schema"].get("properties", {})
       and "sheets" in new_tool["params_schema"]["properties"], list(new_tool["params_schema"].get("properties", {})))

# ---- 3. chạy thử: bảng có chữ, tiền, phần trăm, dòng tổng
args = {
    "ten_tep": f"thu-xuat-{tag}",
    "tieu_de": "Bảng thử xuất Excel",
    "sheets": [{
        "ten": "Chi nhánh",
        "cot": [
            {"nhan": "Chi nhánh", "khoa": "chi_nhanh"},
            {"nhan": "Giải ngân", "khoa": "giai_ngan", "dinh_dang": "tien", "tong": True},
            {"nhan": "Hoàn thành", "khoa": "ty_le", "dinh_dang": "phan_tram"},
            {"nhan": "Số hồ sơ", "khoa": "so_ho_so", "dinh_dang": "so_nguyen"},
        ],
        "dong": [
            {"chi_nhanh": "CN Hà Nội", "giai_ngan": "412.000.000.000", "ty_le": "91,6", "so_ho_so": 34},
            {"chi_nhanh": "=HYPERLINK(\"http://evil\",\"bấm\")", "giai_ngan": 528_500_000_000, "ty_le": 1.057, "so_ho_so": "41"},
            ["CN Đống Đa", 176_300_000_000, 0.8014, 18],
        ],
        "tom_tat": [["Kỳ báo cáo", "09/2026"], ["Người lập", "Kiểm thử"]],
    }],
}
run = c.post(f"/v1/tools/{ex['id']}/test", headers=EB, json={"args": args})
ok("chạy thử trả về tệp", run.status_code == 200 and run.json().get("ok") and run.json()["result"]["tep"],
   run.text[:200])
file = run.json()["result"]["tep"][0] if run.status_code == 200 and run.json().get("result", {}).get("tep") else None
if file:
    created.append(file["id"])
    ok("tệp là .xlsx", file["content_type"].startswith(XLSX_MIME) and file["filename"].endswith(".xlsx"),
       (file["filename"], file["content_type"]))
    dl = c.get(f"/v1/artifacts/{file['id']}/download", headers=EB)
    ok("tải được tệp vừa tạo", dl.status_code == 200 and dl.content[:2] == b"PK", dl.status_code)

    wb, sheet, flat = cells(dl.content)
    values = [x.value for x in flat]
    ok("sheet đúng tên do mô hình đặt", wb.sheetnames == ["Chi nhánh"], wb.sheetnames)
    ok("chuỗi số kiểu Việt Nam thành số thật", 412_000_000_000 in values, values[:14])
    ok("chuỗi số nguyên trong ngoặc kép cũng thành số", 41 in values, values[:16])
    ok("phần trăm ghi 91,6 quy về 0,916",
       any(isinstance(v, float) and abs(v - 0.916) < 1e-6 for v in values),
       [v for v in values if isinstance(v, float)])
    ok("có dòng tổng cộng đúng số",
       "Tổng cộng" in values and any(isinstance(v, (int, float)) and abs(v - 1_116_800_000_000) < 1 for v in values),
       [v for v in values if isinstance(v, (int, float))])
    ok("tóm tắt nằm trên bảng", "Kỳ báo cáo" in values and "09/2026" in values, values[:8])
    sm = c.post(f"/v1/tools/{ex['id']}/test", headers=EB, json={"args": {"ten_tep": f"tom-tat-{tag}", "sheets": [{
        "ten": "Tóm tắt", "cot": [{"nhan": "A", "khoa": "a"}], "dong": [{"a": "x"}],
        "tom_tat": [["Tổng giải ngân", "1.273.400.000.000"], ["Kỳ", "08/2026"], ["Tỷ lệ", "87,2%"]]}]}})
    sm_file = sm.json().get("result", {}).get("tep", [{}])[0] if sm.status_code == 200 and sm.json().get("ok") else None
    if sm_file:
        created.append(sm_file["id"])
        _, _, sm_flat = cells(c.get(f"/v1/artifacts/{sm_file['id']}/download", headers=EB).content)
        sm_values = [x.value for x in sm_flat]
        ok("số trong tóm tắt vào ô số, kỳ và tỷ lệ giữ nguyên chữ",
           1_273_400_000_000 in sm_values and "08/2026" in sm_values and "87,2%" in sm_values, sm_values[:8])
    ok("định dạng tiền và phần trăm được gắn vào ô",
       any(x.number_format.startswith("#,##0") for x in flat) and any(x.number_format == "0.0%" for x in flat))
    ok("ô bắt đầu bằng '=' không thành công thức Excel",
       not any(x.data_type == "f" for x in flat)
       and any(isinstance(v, str) and v.startswith("=HYPERLINK") for v in values),
       [x.data_type for x in flat])

# ---- 3b. tên sheet kiểu mô hình hay đặt (có dấu /) vẫn xuất được
odd = c.post(f"/v1/tools/{ex['id']}/test", headers=EB, json={"args": {
    "ten_tep": f"ky-bao-cao-{tag}",
    "sheets": [
        {"ten": "Doanh số 08/2026", "cot": [{"nhan": "Chi nhánh", "khoa": "cn"}], "dong": [{"cn": "CN Hà Nội"}]},
        {"ten": "Doanh số 08/2026", "cot": [{"nhan": "Nhóm", "khoa": "n"}], "dong": [{"n": "Nhóm 1"}]},
    ]}})
odd_file = odd.json().get("result", {}).get("tep", [{}])[0] if odd.status_code == 200 and odd.json().get("ok") else None
ok("tên sheet có dấu '/' vẫn xuất được (không văng lỗi)", odd_file is not None, odd.text[:200])
if odd_file:
    created.append(odd_file["id"])
    wb2 = load_workbook(io.BytesIO(c.get(f"/v1/artifacts/{odd_file['id']}/download", headers=EB).content))
    ok("tên sheet được đổi cho hợp lệ và không trùng nhau",
       wb2.sheetnames == ["Doanh số 08-2026", "Doanh số 08-2026 (2)"], wb2.sheetnames)

# ---- 4. bảng rỗng thì báo lỗi rõ ràng, không tạo tệp rác
empty = c.post(f"/v1/tools/{ex['id']}/test", headers=EB,
               json={"args": {"sheets": [{"ten": "Trống", "cot": [{"nhan": "A", "khoa": "a"}], "dong": []}]}})
ok("bảng không có dòng nào bị từ chối", empty.status_code >= 400 or not empty.json().get("ok"), empty.text[:200])
missing = c.post(f"/v1/tools/{ex['id']}/test", headers=EB, json={"args": {"tieu_de": "Thiếu sheets"}})
ok("thiếu 'sheets' bị chặn theo schema", missing.status_code >= 400 or not missing.json().get("ok"), missing.text[:160])

# ---- 5. trong chat: cán bộ đưa dữ liệu rồi bảo xuất Excel
apps = {a["name"]: a for a in c.get("/v1/audit/filters", headers=A).json()["apps"]}
app = apps["Trợ lý Hồ sơ Tín dụng"]
st = c.post("/v1/staff/login", json={"email": "rm.an@msb-demo.vn", "password": "demo123"}).json()["access_token"]
S = {"Authorization": f"Bearer {st}"}
question = ("Xuất giúp tôi ra Excel bảng theo dõi sau: HS2026-0412 khách Thép Đông Á 12 tỷ quá hạn 3 ngày; "
            "HS2026-0518 khách Nhựa Tiền Phong 8,5 tỷ quá hạn 7 ngày; HS2026-0777 khách Dệt May Sài Gòn "
            "1,25 tỷ quá hạn 1 ngày. Cột: mã hồ sơ, khách hàng, số tiền, số ngày quá hạn.")
with c.stream("POST", f"/v1/staff/apps/{app['id']}/chat", headers=S, json={"message": question}) as r:
    ev = sse(r)
called = [e.get("tool") for e in ev if e["type"] == "tool_call"]
files = [e for e in ev if e["type"] == "artifact"]
ok("agent gọi công cụ xuất Excel", "xuat_excel" in called, called)
ok("chat trả về tệp .xlsx", any((f.get("content_type") or "").startswith(XLSX_MIME) for f in files),
   [(f.get("filename"), f.get("content_type")) for f in files])
chat_file = next((f for f in files if (f.get("content_type") or "").startswith(XLSX_MIME)), None)
answer = "".join(e["content"] for e in ev if e["type"] == "token")
ok("câu trả lời ngắn gọn, không chép lại cả bảng", 10 < len(answer) < 1500, len(answer))

if chat_file:
    dl = c.get(f"/v1/staff/reports/{chat_file['artifact_id']}/download", headers=S)
    ok("cán bộ tải được tệp ngay trong chat", dl.status_code == 200 and dl.content[:2] == b"PK", dl.status_code)
    if dl.status_code == 200:
        _, _, flat = cells(dl.content)
        values = [x.value for x in flat]
        ok("bảng chứa đúng dữ liệu của hội thoại",
           any(isinstance(v, str) and "HS2026-0412" in v for v in values)
           and any(isinstance(v, str) and "HS2026-0777" in v for v in values), values[:12])
        ok("số tiền vào ô dạng số, không phải chữ",
           any(isinstance(v, (int, float)) and not isinstance(v, bool) and v >= 1000 for v in values),
           [v for v in values if isinstance(v, (int, float))])

    # ---- 6. tệp thuộc về chính cán bộ đó
    inbox = c.get("/v1/staff/reports", headers=S).json()
    ok("tệp nằm trong 'Báo cáo của tôi' của người hỏi",
       any(x["id"] == chat_file["artifact_id"] for x in inbox), [x["title"] for x in inbox[:3]])
    other = c.post("/v1/staff/login", json={"email": "ca.binh@msb-demo.vn", "password": "demo123"}).json()["access_token"]
    theirs = c.get("/v1/staff/reports", headers={"Authorization": f"Bearer {other}"}).json()
    ok("cán bộ khác không thấy tệp này", not any(x["id"] == chat_file["artifact_id"] for x in theirs), len(theirs))
    created.append(chat_file["artifact_id"])

# ---- 6b. gắn công cụ mà quên bật công tắc agent thì trợ lý không gọi gì cả
quiet = c.post("/v1/apps", headers=EB, json={
    "name": f"Trợ lý quên bật {tag}", "dataset_ids": [], "tool_ids": [ex["id"]], "agent_enabled": False}).json()
try:
    with c.stream("POST", f"/v1/apps/{quiet['id']}/test-chat", headers=EB,
                  json={"message": "Xuất giúp tôi ra Excel bảng: A 1; B 2. Cột: tên, số."}) as r:
        ev = sse(r)
    ok("công tắc tắt: có gắn công cụ nhưng không gọi, không sinh tệp",
       not [e for e in ev if e["type"] == "artifact"] and not [e for e in ev if e["type"] == "tool_call"],
       [e.get("tool") for e in ev if e["type"] == "tool_call"])

    c.patch(f"/v1/apps/{quiet['id']}", headers=EB, json={"agent_enabled": True})
    with c.stream("POST", f"/v1/apps/{quiet['id']}/test-chat", headers=EB,
                  json={"message": "Xuất giúp tôi ra Excel bảng: A 1; B 2. Cột: tên, số."}) as r:
        ev = sse(r)
    admin_files = [e for e in ev if e["type"] == "artifact"]
    ok("bật công tắc lên thì xuất được ngay", bool(admin_files),
       [e.get("tool") for e in ev if e["type"] == "tool_call"])
    for f in admin_files:
        c.delete(f"/v1/artifacts/{f['artifact_id']}", headers=EB)
finally:
    c.delete(f"/v1/apps/{quiet['id']}", headers=EB)

# ---- 7. chào hỏi vẫn không sinh tệp
with c.stream("POST", f"/v1/staff/apps/{app['id']}/chat", headers=S, json={"message": "xin chào"}) as r:
    ev = sse(r)
ok("chào hỏi không gọi công cụ xuất tệp",
   not any(e.get("tool") == "xuat_excel" for e in ev if e["type"] == "tool_call")
   and not [e for e in ev if e["type"] == "artifact"],
   [e.get("tool") for e in ev if e["type"] == "tool_call"])

# ---- dọn dẹp
for artifact_id in created:
    c.delete(f"/v1/artifacts/{artifact_id}", headers=EB)
if new_tool:
    c.delete(f"/v1/tools/{new_tool['id']}", headers=EB)
ok("dọn công cụ và tệp thử", all(c.get(f"/v1/artifacts/{a}", headers=EB).status_code == 404 for a in created))

passed = sum(1 for _, v in results if v)
print(f"\n===== EXPORT: {passed}/{len(results)} passed =====")
sys.exit(0 if passed == len(results) else 1)
