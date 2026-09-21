"""Biểu mẫu nghiệp vụ: điền sẵn từ hệ thống lõi, gợi ý AI cho trường tự luận (không lộ PII),
xuất .docx, và phạm vi đơn vị.

Cần: API :8000, mock core :8095, provider LLM. Tự dọn tệp đã tạo.
Chạy:  apps/api/.venv/bin/python scratch/smoke_forms.py
       API=https://<domain> ADMIN_PASSWORD=... apps/api/.venv/bin/python scratch/smoke_forms.py
"""
import io
import os
import sys
import zipfile

import httpx

API = os.environ.get("API", "http://localhost:8000")
c = httpx.Client(base_url=API, timeout=180)
results = []


def ok(label, cond, extra=""):
    results.append((label, bool(cond)))
    print(("PASS " if cond else "FAIL ") + label, "" if cond else "→ " + str(extra)[:300])
    return bool(cond)


tok = c.post("/v1/auth/login", json={"email": "admin@querion.io",
                                     "password": os.environ.get("ADMIN_PASSWORD", "admin123")}).json()["access_token"]
A = {"Authorization": f"Bearer {tok}"}
ws = {w["name"]: w["id"] for w in c.get("/v1/audit/filters", headers=A).json()["workspaces"]}
EB = {**A, "X-Workspace-Id": ws["Khối Khách hàng Doanh nghiệp (EB)"]}
OPS = {**A, "X-Workspace-Id": ws["Khối Vận hành & Thanh toán quốc tế"]}

forms = c.get("/v1/forms", headers=EB).json()
form = next((f for f in forms if f["name"] == "Đề nghị giải ngân khoản vay"), None)
if not ok("có biểu mẫu demo sau khi seed", form is not None, [f["name"] for f in forms]):
    sys.exit(1)
ok("biểu mẫu đã công bố và có mẫu .docx", form["is_published"] and form["has_template"], form)
ok("biểu mẫu khai báo đủ 3 nguồn dữ liệu (cán bộ nhập / công cụ / AI)",
   {f["source"] for f in form["fields"]} == {"user", "tool", "llm"}, {f["source"] for f in form["fields"]})
ok("trường tên khách hàng được đánh dấu nhạy cảm (PII)",
   next(f for f in form["fields"] if f["name"] == "khach_hang")["pii"] is True)

# ---- cán bộ dùng biểu mẫu
st = c.post("/v1/staff/login", json={"email": "rm.an@msb-demo.vn", "password": "demo123"}).json()["access_token"]
S = {"Authorization": f"Bearer {st}"}
listed = c.get("/v1/staff/forms", headers=S).json()
ok("cán bộ EB thấy biểu mẫu", any(f["id"] == form["id"] for f in listed), [f["name"] for f in listed])
mine = next(f for f in listed if f["id"] == form["id"])
ok("biểu mẫu báo có chức năng điền sẵn", mine["has_prefill"] and mine["prefill_label"] == "Mã hồ sơ", mine)

st_ops = c.post("/v1/staff/login", json={"email": "gdv.cuong@msb-demo.vn", "password": "demo123"}).json()["access_token"]
ok("cán bộ đơn vị khác không thấy biểu mẫu của EB",
   not any(f["id"] == form["id"] for f in c.get("/v1/staff/forms", headers={"Authorization": f"Bearer {st_ops}"}).json()))

# ---- điền sẵn từ hệ thống lõi
r = c.post(f"/v1/staff/forms/{form['id']}/prefill", headers=S, json={"key": "HS2026-0412"})
ok("điền sẵn từ mã hồ sơ: 200", r.status_code == 200, r.text[:200])
values = r.json().get("values", {})
ok("điền đúng khách hàng, sản phẩm, số tiền từ hệ thống lõi",
   values.get("khach_hang") == "Công ty CP Thép Đông Á" and values.get("so_tien") == 12_000_000_000
   and "Vay bổ sung" in str(values.get("san_pham")), values)
miss = c.post(f"/v1/staff/forms/{form['id']}/prefill", headers=S, json={"key": "HS-KHONG-CO"})
ok("mã không tồn tại báo lỗi rõ ràng (400/404)", miss.status_code in (400, 404), miss.text[:160])
deny = c.post(f"/v1/staff/forms/{form['id']}/prefill", headers={"Authorization": f"Bearer {st_ops}"},
              json={"key": "HS2026-0412"})
ok("đơn vị khác không gọi được điền sẵn (404)", deny.status_code == 404, deny.status_code)

# ---- gợi ý AI cho trường tự luận, không được thấy dữ liệu nhạy cảm
filled = {**values, "ma_ho_so": "HS2026-0412", "so_tien_giai_ngan": 4_000_000_000,
          "muc_dich": "Thanh toán tiền mua thép cuộn cán nóng theo hợp đồng số 18/2026/HĐMB",
          "ngay_de_nghi": "18/09/2026"}
r = c.post(f"/v1/staff/forms/{form['id']}/suggest", headers=S, json={"field": "nhan_xet", "values": filled})
ok("gợi ý AI cho phần nhận xét: 200", r.status_code == 200, r.text[:200])
suggestion = r.json().get("text", "")
ok("nội dung gợi ý đủ dài và bằng tiếng Việt", len(suggestion) > 60, suggestion[:160])
ok("gợi ý KHÔNG nhắc tên khách hàng (trường PII không vào prompt)",
   "Thép Đông Á" not in suggestion, suggestion[:300])
bad_field = c.post(f"/v1/staff/forms/{form['id']}/suggest", headers=S, json={"field": "khach_hang", "values": filled})
ok("không xin gợi ý cho trường không phải AI (400)", bad_field.status_code == 400, bad_field.text[:160])

# ---- thiếu trường bắt buộc
missing = c.post(f"/v1/staff/forms/{form['id']}/submit", headers=S,
                 json={"values": {k: v for k, v in filled.items() if k != "muc_dich"}})
ok("thiếu trường bắt buộc bị chặn kèm tên trường (400)",
   missing.status_code == 400 and "Mục đích" in missing.text, missing.text[:200])

# ---- xuất .docx
submit = c.post(f"/v1/staff/forms/{form['id']}/submit", headers=S,
                json={"values": {**filled, "nhan_xet": suggestion}})
ok("xuất biểu mẫu: 201 + artifact", submit.status_code == 201 and submit.json().get("artifact_id"), submit.text[:200])
artifact_id = submit.json().get("artifact_id")

dl = c.get(f"/v1/staff/reports/{artifact_id}/download", headers=S)
ok("cán bộ tải được biểu mẫu vừa lập", dl.status_code == 200 and dl.content[:2] == b"PK", dl.status_code)
with zipfile.ZipFile(io.BytesIO(dl.content)) as z:
    xml = z.read("word/document.xml").decode("utf-8")
ok("file Word đã điền dữ liệu, không còn placeholder",
   "HS2026-0412" in xml and "{{" not in xml and "{%" not in xml, [("HS2026-0412" in xml), ("{{" in xml)])
ok("file có tên khách hàng và số tiền định dạng VN",
   "Thép Đông Á" in xml and "4.000.000.000" in xml, [("Thép Đông Á" in xml), ("4.000.000.000" in xml)])
ok("file ghi tên cán bộ lập", "Nguyễn Văn An" in xml)

# ---- nhật ký + token
runs = c.get("/v1/audit/runs", headers=A, params={"days": 1, "channel": "form"}).json()
form_runs = [r for r in runs if r.get("channel") == "form"]
ok("nhật ký ghi lượt lập biểu mẫu (channel=form)", len(form_runs) >= 2, len(form_runs))
usage = c.get("/v1/usage/summary", headers=A, params={"days": 1}).json()
ok("token của gợi ý biểu mẫu được tính riêng (form_suggest)",
   any(b["key"] == "form_suggest" for b in usage["by_component"]), [b["key"] for b in usage["by_component"]])

# ---- quản trị: mẫu .docx phải khớp trường đã khai báo
bad_tpl = c.post(f"/v1/forms/{form['id']}/template", headers=EB,
                 files={"file": ("x.txt", b"khong phai docx", "text/plain")})
ok("từ chối mẫu không phải .docx (400)", bad_tpl.status_code == 400, bad_tpl.text[:120])
r = c.post("/v1/forms", headers=EB, json={"name": "Biểu mẫu kiểm thử", "fields": [
    {"name": "so xau", "label": "Tên sai", "type": "string"}]})
ok("tên trường không hợp lệ bị chặn (400)", r.status_code == 400, r.text[:160])
tmp = c.post("/v1/forms", headers=EB, json={"name": "Biểu mẫu kiểm thử", "fields": [
    {"name": "ghi_chu", "label": "Ghi chú", "type": "string"}]}).json()
pub = c.patch(f"/v1/forms/{tmp['id']}", headers=EB, json={"is_published": True})
ok("không công bố được biểu mẫu chưa có mẫu .docx (400)", pub.status_code == 400, pub.text[:160])
def _docx_with(text: str) -> bytes:
    from docx import Document

    doc, buf = Document(), io.BytesIO()
    doc.add_paragraph(text)
    doc.save(buf)
    return buf.getvalue()


DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
wrong_tpl = c.post(f"/v1/forms/{tmp['id']}/template", headers=EB,
                   files={"file": ("mau.docx", _docx_with("Hồ sơ {{ ma_ho_so }} — ghi chú {{ ghi_chu }}"), DOCX_MIME)})
ok("mẫu dùng biến chưa khai báo bị chặn kèm tên biến",
   wrong_tpl.status_code == 400 and "ma_ho_so" in wrong_tpl.text, wrong_tpl.text[:200])
good_tpl = c.post(f"/v1/forms/{tmp['id']}/template", headers=EB,
                  files={"file": ("mau.docx", _docx_with("Ghi chú: {{ ghi_chu }} — lập {{ today }}"), DOCX_MIME)})
ok("mẫu chỉ dùng biến đã khai báo được chấp nhận", good_tpl.status_code == 200 and good_tpl.json()["has_template"],
   good_tpl.text[:160])
ok("có mẫu rồi thì công bố được",
   c.patch(f"/v1/forms/{tmp['id']}", headers=EB, json={"is_published": True}).json().get("is_published") is True)
ok("đơn vị khác không sửa được biểu mẫu (404)",
   c.patch(f"/v1/forms/{tmp['id']}", headers=OPS, json={"name": "đổi trộm"}).status_code == 404)
c.delete(f"/v1/forms/{tmp['id']}", headers=EB)
ok("xoá biểu mẫu tạm", all(f["id"] != tmp["id"] for f in c.get("/v1/forms", headers=EB).json()))

c.delete(f"/v1/artifacts/{artifact_id}", headers=EB)
ok("dọn tệp biểu mẫu thử", c.get(f"/v1/artifacts/{artifact_id}", headers=EB).status_code == 404)

passed = sum(1 for _, v in results if v)
print(f"\n===== FORMS: {passed}/{len(results)} passed =====")
sys.exit(0 if passed == len(results) else 1)
