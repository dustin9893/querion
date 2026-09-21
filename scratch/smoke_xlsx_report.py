"""Báo cáo Excel: workflow gọi MCP kho dữ liệu, xuất .xlsx nhiều sheet + tóm tắt Markdown.

Cần: API :8000, jobs worker, mock DWH :8097 (./scripts/mock-dwh.sh), provider LLM.
Chạy:  apps/api/.venv/bin/python scratch/smoke_xlsx_report.py
       API=https://<domain> ADMIN_PASSWORD=... apps/api/.venv/bin/python scratch/smoke_xlsx_report.py
"""
import io
import os
import sys
import time

import httpx

API = os.environ.get("API", "http://localhost:8000")
c = httpx.Client(base_url=API, timeout=240)
results = []
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def ok(label, cond, extra=""):
    results.append((label, bool(cond)))
    print(("PASS " if cond else "FAIL ") + label, "" if cond else "→ " + str(extra)[:300])
    return bool(cond)


tok = c.post("/v1/auth/login", json={"email": "admin@querion.io",
                                     "password": os.environ.get("ADMIN_PASSWORD", "admin123")}).json()["access_token"]
A = {"Authorization": f"Bearer {tok}"}
ws = {w["name"]: w["id"] for w in c.get("/v1/audit/filters", headers=A).json()["workspaces"]}
EB = {**A, "X-Workspace-Id": ws["Khối Khách hàng Doanh nghiệp (EB)"]}

wf = next((w for w in c.get("/v1/workflows", headers=EB).json() if w["name"] == "Báo cáo kinh doanh tháng (Excel)"), None)
if not ok("có luồng báo cáo Excel sau khi seed", wf is not None):
    sys.exit(1)
nodes = {n["id"]: n for n in wf["graph_json"]["nodes"]}
ok("dùng công cụ MCP (3 lần gọi kho dữ liệu)",
   sum(1 for n in nodes.values() if n["type"] == "tool_call" and n["data"].get("mcp_tool")) == 3,
   [n["data"].get("mcp_tool") for n in nodes.values() if n["type"] == "tool_call"])
ok("node Excel khai báo 3 sheet", len(nodes["xuat_xlsx"]["data"]["sheets"]) == 3,
   [s["name"] for s in nodes["xuat_xlsx"]["data"].get("sheets", [])])

# ---- chạy nền
r = c.post(f"/v1/workflows/{wf['id']}/jobs", headers=EB, json={"query": "bao cao", "inputs": {"thang": "09/2026"}})
ok("xếp hàng chạy: 202", r.status_code == 202, r.text[:200])
run_id = r.json()["run_id"]
row = {}
for _ in range(60):
    time.sleep(4)
    row = next((x for x in c.get(f"/v1/workflows/{wf['id']}/runs", headers=EB).json() if x["id"] == run_id), {})
    if row.get("status") in ("completed", "failed"):
        break
ok("chạy xong completed", row.get("status") == "completed", row)
arts = {a["content_type"].split(";")[0]: a for a in row.get("artifacts", [])}
ok("sinh cả Excel và Markdown", XLSX_MIME in arts and "text/markdown" in arts, list(arts))

# ---- nội dung Excel
xlsx = arts.get(XLSX_MIME)
dl = c.get(f"/v1/artifacts/{xlsx['id']}/download", headers=EB)
ok("tải Excel: đúng content-type", dl.status_code == 200 and dl.headers["content-type"].startswith(XLSX_MIME), dl.headers.get("content-type"))
from openpyxl import load_workbook  # noqa: E402

wb = load_workbook(io.BytesIO(dl.content))
ok("workbook có 3 sheet đúng tên", wb.sheetnames == ["Doanh số chi nhánh", "Nợ theo nhóm", "KPI cán bộ"], wb.sheetnames)

sheet = wb["Doanh số chi nhánh"]
values = [[cell.value for cell in row] for row in sheet.iter_rows()]
flat = [v for row in values for v in row if v is not None]
ok("sheet doanh số có dữ liệu từ kho dữ liệu (MCP)", "CN Hồ Chí Minh" in flat and 528_500_000_000 in flat,
   [v for v in flat[:12]])
ok("có dòng tổng cộng", "Tổng cộng" in flat, flat[-8:])
money_cells = [cell for row in sheet.iter_rows() for cell in row if cell.number_format.startswith("#,##0")]
ok("số tiền được định dạng số (không phải chữ)", len(money_cells) >= 10, len(money_cells))
pct = [cell for row in sheet.iter_rows() for cell in row if cell.number_format == "0.0%"]
ok("tỷ lệ hoàn thành ở dạng phần trăm", len(pct) >= 5, len(pct))
ok("cố định dòng tiêu đề để cuộn bảng", sheet.freeze_panes is not None, sheet.freeze_panes)

kpi = wb["KPI cán bộ"]
kpi_flat = [c.value for row in kpi.iter_rows() for c in row if c.value is not None]
ok("sheet KPI có cán bộ và tỷ lệ đúng hạn", "MSB01001" in kpi_flat and any(isinstance(v, float) and 0 < v < 1 for v in kpi_flat),
   kpi_flat[:10])

# ---- tóm tắt Markdown
md = arts["text/markdown"]
preview = c.get(f"/v1/artifacts/{md['id']}", headers=EB).json().get("preview", "")
ok("tóm tắt Markdown có số liệu tổng hợp", "Tổng giải ngân" in preview and "đ" in preview, preview[:200])
ok("tóm tắt có phần nhận định của LLM", "## Nhận định" in preview and len(preview.split("## Nhận định")[1]) > 80, preview[-300:])

# ---- tham số tháng khác cho ra số khác
r2 = c.post(f"/v1/workflows/{wf['id']}/jobs", headers=EB, json={"query": "x", "inputs": {"thang": "08/2026"}})
run2 = r2.json()["run_id"]
row2 = {}
for _ in range(60):
    time.sleep(4)
    row2 = next((x for x in c.get(f"/v1/workflows/{wf['id']}/runs", headers=EB).json() if x["id"] == run2), {})
    if row2.get("status") in ("completed", "failed"):
        break
x2 = next((a for a in row2.get("artifacts", []) if a["content_type"].startswith(XLSX_MIME)), None)
if ok("chạy kỳ 08/2026 cũng ra Excel", x2 is not None, row2.get("status")):
    wb2 = load_workbook(io.BytesIO(c.get(f"/v1/artifacts/{x2['id']}/download", headers=EB).content))
    flat2 = [v for row in wb2["Doanh số chi nhánh"].iter_rows() for v in (c2.value for c2 in row) if v is not None]
    ok("số liệu đổi theo kỳ báo cáo", 474_000_000_000 in flat2 and 528_500_000_000 not in flat2, flat2[:10])

# ---- nhật ký
audit = c.get(f"/v1/audit/runs/{run_id}", headers=A).json()
steps = [s["node_type"] for s in audit.get("steps", [])]
ok("nhật ký ghi 3 bước gọi công cụ + 2 bước xuất tài liệu",
   steps.count("tool_call") == 3 and steps.count("render_document") == 2, steps)

# ---- dọn dẹp
for a in list(arts.values()) + ([x2] if x2 else []):
    c.delete(f"/v1/artifacts/{a['id']}", headers=EB)
for a in row2.get("artifacts", []):
    c.delete(f"/v1/artifacts/{a['id']}", headers=EB)
ok("dọn tệp thử", c.get(f"/v1/artifacts/{xlsx['id']}", headers=EB).status_code == 404)

passed = sum(1 for _, v in results if v)
print(f"\n===== XLSX REPORT: {passed}/{len(results)} passed =====")
sys.exit(0 if passed == len(results) else 1)
