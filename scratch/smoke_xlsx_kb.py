"""Tệp .xlsx trong kho tri thức: tải lên, lập chỉ mục theo dòng, truy vấn, trích dẫn, bật / tắt / xoá.

Chứng minh: .xlsx (kể cả khai báo sai content type, đuôi viết hoa) được nhận, .xls / .csv bị từ chối;
bảng tính ra các đoạn "[Sheet <tên> › dòng a–b] ..." — ô gộp dọc lặp trên từng dòng, ngày / phần trăm /
số tiền hiển thị kiểu Việt Nam, "<" đổi thành "＜", sheet / dòng / cột ẩn không lọt vào đoạn nào;
bảng tính rỗng và tệp giả mạo báo lỗi tiếng Việt thay vì kẹt "indexing"; truy vấn và trợ lý trích dẫn
đúng sheet; lập chỉ mục lại cho ra đúng các đoạn cũ; .txt / .docx / .pdf vẫn chạy như trước.

Cần: API, worker ĐÃ BUILD LẠI image có openpyxl, provider embedding + LLM đang bật.
Tự tạo và tự xoá kho + trợ lý tạm.
Chạy:  API=http://127.0.0.1:8001 apps/api/.venv/Scripts/python.exe scratch/smoke_xlsx_kb.py
       SKIP_CHAT=1 ...   # bỏ vòng hỏi trợ lý (chậm)
"""
import datetime as dt
import io
import json
import os
import re
import sys
import time
import uuid

import httpx

API = os.environ.get("API", "http://localhost:8000")
SKIP_CHAT = os.environ.get("SKIP_CHAT") == "1"
c = httpx.Client(base_url=API, timeout=180)
results = []

XLSX_CT = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
DOCX_CT = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
SECTION_RE = re.compile(r"^\[([^\]]{1,160})\]\s")


def ok(label, cond, extra=""):
    results.append((label, bool(cond)))
    print(("PASS " if cond else "FAIL ") + label, ("" if cond else "→ ") + str(extra)[:300])
    return bool(cond)


def sse(resp):
    out = []
    for line in resp.iter_lines():
        if line.startswith("data: ") and line[6:] != "[DONE]":
            try:
                out.append(json.loads(line[6:]))
            except Exception:
                pass
    return out


def wait_ready(doc_id, H, timeout=180):
    end = time.time() + timeout
    while time.time() < end:
        d = c.get(f"/v1/documents/{doc_id}", headers=H).json()
        if d["status"] in ("ready", "failed"):
            return d
        time.sleep(2)
    return d


def all_chunks(doc_id, H):
    out, page = [], 1
    while True:
        body = c.get(f"/v1/documents/{doc_id}/chunks", headers=H, params={"page": page, "page_size": 100}).json()
        out += [ch["content"] for ch in body["chunks"]]
        if not body["chunks"] or len(out) >= body["total"]:
            return out
        page += 1


# ---------------------------------------------------------------------------
# Workbooks and control files, built in memory (synthetic content only)
# ---------------------------------------------------------------------------

GROUPS = ["Tiền gửi", "Chuyển tiền", "Thẻ", "Tín dụng"]
ROWS_PER_GROUP = 20


def build_main_xlsx(tag):
    """Fee table with the traps the reader must handle. Returns (bytes, facts for the checks)."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Biểu phí"
    ws["A1"] = "BIỂU PHÍ DỊCH VỤ THỬ NGHIỆM"
    ws.merge_cells("A1:F1")
    for col, name in enumerate(["Nhóm", "Dịch vụ", "Mức giao dịch", "Mức phí", "Tối thiểu", "Ngày hiệu lực", "Ghi chú nội bộ"], 1):
        ws.cell(row=2, column=col, value=name)

    facts = {
        "hidden": [f"BIMAT-SHEET-{tag}", f"BIMAT-DONG-{tag}", f"BIMAT-COT-{tag}"],
        "service": f"Phí xác nhận số dư mã {tag}",
        "group_rows": {},  # group name -> number of data rows
    }
    row = 3
    for g, group in enumerate(GROUPS, 1):
        first = row
        label = f"Nhóm {g} — {group}"
        for i in range(1, ROWS_PER_GROUP + 1):
            service = f"Dịch vụ G{g}-{i:02d} thuộc nhóm {group.lower()} áp dụng cho khách hàng doanh nghiệp"
            ws.cell(row=row, column=2, value=service)
            ws.cell(row=row, column=3, value=f"Từ {i * 10} triệu đến {i * 10 + 10} triệu đồng")
            fee = ws.cell(row=row, column=4, value=1000 * (g * 100 + i))
            fee.number_format = "#,##0"
            ws.cell(row=row, column=5, value=f"{5000 + i * 100} đồng / giao dịch")
            eff = ws.cell(row=row, column=6, value=dt.datetime(2026, 1, 1))
            eff.number_format = "dd/mm/yyyy"
            ws.cell(row=row, column=7, value=f"BIMAT-COT-{tag}")
            row += 1
        ws.cell(row=first, column=1, value=label)
        ws.merge_cells(start_row=first, start_column=1, end_row=row - 1, end_column=1)
        facts["group_rows"][label] = ROWS_PER_GROUP

    # The tagged row: unique amount, percent fee, "<" in the band, a real date.
    ws.cell(row=row, column=1, value="Nhóm 5 — Xác nhận")
    ws.cell(row=row, column=2, value=facts["service"])
    ws.cell(row=row, column=3, value="<500 triệu")
    amount = ws.cell(row=row, column=4, value=73412)
    amount.number_format = "#,##0"
    pct = ws.cell(row=row, column=5, value=0.0002)
    pct.number_format = "0.00%"
    ws.cell(row=row, column=6, value=dt.date(2026, 3, 1))
    ws.cell(row=row, column=7, value=f"BIMAT-COT-{tag}")
    row += 1

    # A hidden row inside the table.
    for col in range(1, 7):
        ws.cell(row=row, column=col, value=f"BIMAT-DONG-{tag}" if col == 2 else f"ẩn {col}")
    ws.row_dimensions[row].hidden = True
    row += 1
    ws.cell(row=row, column=1, value="Nhóm 6 — Khác")
    ws.cell(row=row, column=2, value="Dịch vụ cuối bảng sau dòng ẩn")
    ws.cell(row=row, column=4, value=15000).number_format = "#,##0"

    ws.column_dimensions["G"].hidden = True

    hidden = wb.create_sheet("Nội bộ")
    hidden["A1"] = "Khoá"
    hidden["B1"] = "Giá trị"
    hidden["C1"] = "Ghi chú"
    hidden["A2"] = "k1"
    hidden["B2"] = f"BIMAT-SHEET-{tag}"
    hidden["C2"] = "không được lập chỉ mục"
    hidden.sheet_state = "hidden"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), facts


def build_empty_xlsx():
    import openpyxl

    buf = io.BytesIO()
    openpyxl.Workbook().save(buf)
    return buf.getvalue()


def build_docx(lines):
    import docx

    d = docx.Document()
    for line in lines:
        d.add_paragraph(line)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def build_pdf(text):
    """One-page PDF with a single line of ASCII text (Helvetica), xref offsets computed."""
    safe = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = f"BT /F1 12 Tf 72 720 Td ({safe}) Tj ET".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for n, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += f"{n} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{o:010d} 00000 n \n".encode() for o in offsets)
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return bytes(out)


def build_all(tag):
    main, facts = build_main_xlsx(tag)
    txt_tag, docx_tag, pdf_tag = f"TXT{tag}", f"DOCX{tag}", f"PDF{tag}"
    files = {
        "main": (f"Biểu phí thử nghiệm {tag}.xlsx", main, XLSX_CT),
        "upper": (f"Biểu phí thử nghiệm {tag} bản sao.XLSX", main, "text/plain"),
        "txt": (f"doi_chung_{tag}.txt", (
            f"Điều 1. Giờ làm việc của quầy thử nghiệm {txt_tag}\n"
            f"Quầy giao dịch thử nghiệm {txt_tag} mở cửa từ 8 giờ đến 17 giờ các ngày làm việc trong tuần.\n"
        ).encode(), "text/plain"),
        "docx": (f"doi_chung_{tag}.docx", build_docx([
            f"Quy định thử nghiệm {docx_tag}",
            f"Hồ sơ thử nghiệm {docx_tag} được lưu trữ trong 10 năm kể từ ngày tất toán.",
        ]), DOCX_CT),
        "pdf": (f"doi_chung_{tag}.pdf", build_pdf(
            f"Test control document {pdf_tag}: the branch opens at 8 am on working days."), "application/pdf"),
        "empty": (f"rong_{tag}.xlsx", build_empty_xlsx(), XLSX_CT),
        "fake": ("gia-mao.xlsx", build_docx([f"Văn bản Word giả mạo bảng tính {tag}"]), XLSX_CT),
    }
    tags = {"txt": txt_tag, "docx": docx_tag, "pdf": pdf_tag}
    return files, facts, tags


if __name__ == "__main__" and os.environ.get("BUILD_ONLY") == "1":
    # Offline self-check of the builders: python scratch/smoke_xlsx_kb.py with BUILD_ONLY=1
    files, facts, tags = build_all("ABC123")
    for k, (name, data, ct) in files.items():
        print(f"{k:6} {name!r:50} {len(data):7} bytes {ct}")
    print(facts, tags)
    sys.exit(0)


# ---------------------------------------------------------------------------
# Live run
# ---------------------------------------------------------------------------

tok = c.post("/v1/auth/login", json={"email": "admin@querion.io", "password": os.environ.get("ADMIN_PASSWORD", "admin123")}).json()["access_token"]
A = {"Authorization": f"Bearer {tok}"}
ws = {w["name"]: w["id"] for w in c.get("/v1/audit/filters", headers=A).json()["workspaces"]}
EB = {**A, "X-Workspace-Id": ws["Khối Khách hàng Doanh nghiệp (EB)"]}
tag = uuid.uuid4().hex[:6].upper()
FILES, FACTS, TAGS = build_all(tag)
QUESTION = f"{FACTS['service']} có mức phí là bao nhiêu?"
TXT_QUESTION = f"Quầy giao dịch thử nghiệm {TAGS['txt']} mở cửa mấy giờ?"

ds = c.post("/v1/datasets", headers=EB, json={"name": f"Kiểm thử xlsx {tag}", "visibility": "internal"}).json()
app = None


def upload(name, data, ct):
    return c.post(f"/v1/datasets/{ds['id']}/documents/upload", headers=EB,
                  files={"file": (name, data, ct)}, data={"doc_type": "bieu_phi"})


def retrieved(query, top_k=10):
    r = c.post("/v1/retrieval", headers=EB, json={"query": query, "dataset_ids": [ds["id"]], "top_k": top_k})
    return r.json()["results"]


def ids(hits):
    return {h["document_id"] for h in hits}


try:
    # ---- upload: accepted and rejected types
    r_main = upload(*FILES["main"])
    ok("tải .xlsx lên: 201", r_main.status_code == 201, r_main.text)
    r_upper = upload(*FILES["upper"])
    ok("tải .XLSX khai báo text/plain: 201", r_upper.status_code == 201, r_upper.text)
    r = upload(f"cu_{tag}.xls", b"\xd0\xcf\x11\xe0 not really xls", "application/vnd.ms-excel")
    ok("tệp .xls bị từ chối (400, liệt kê định dạng cho phép)",
       r.status_code == 400 and "Allowed: .docx, .pdf, .txt, .xlsx" in r.json().get("detail", ""), r.text)
    r = upload(f"bang_{tag}.csv", b"a,b,c\n1,2,3\n", "text/csv")
    ok("tệp .csv bị từ chối (400)", r.status_code == 400, r.status_code)
    ctrl = {k: upload(*FILES[k]) for k in ("txt", "docx", "pdf")}
    ok("tải tệp đối chứng .txt / .docx / .pdf: 201", all(x.status_code == 201 for x in ctrl.values()),
       {k: x.status_code for k, x in ctrl.items()})
    main_id, upper_id = r_main.json()["id"], r_upper.json()["id"]
    ctrl_ids = {k: x.json()["id"] for k, x in ctrl.items()}

    # ---- indexing
    d = wait_ready(main_id, EB)
    if not ok("bảng tính lập chỉ mục xong, ≥ 10 đoạn", d["status"] == "ready" and (d.get("chunk_count") or 0) >= 10,
              f"{d['status']} chunk_count={d.get('chunk_count')} error={d.get('error_message')} "
              "— worker image chưa build lại với openpyxl?"):
        sys.exit(1)
    main_count = d["chunk_count"]
    d = wait_ready(upper_id, EB)
    upper_chunks = all_chunks(upper_id, EB) if d["status"] == "ready" else []
    ok(".XLSX khai báo text/plain vẫn đọc như bảng tính (đoạn đầu '[Sheet ')",
       upper_chunks and upper_chunks[0].startswith("[Sheet "), upper_chunks[:1] or d)
    c.delete(f"/v1/documents/{upper_id}", headers=EB)  # an identical copy would crowd the retrieval checks

    ctrl_docs = {k: wait_ready(i, EB) for k, i in ctrl_ids.items()}
    ctrl_chunks = {k: all_chunks(i, EB) if ctrl_docs[k]["status"] == "ready" else [] for k, i in ctrl_ids.items()}
    ok("đối chứng sẵn sàng: txt bắt đầu '[Điều 1', docx / pdf chứa mã của mình",
       all(x["status"] == "ready" for x in ctrl_docs.values())
       and ctrl_chunks["txt"] and ctrl_chunks["txt"][0].startswith("[Điều 1")
       and any(TAGS["docx"] in ch for ch in ctrl_chunks["docx"])
       and any(TAGS["pdf"] in ch for ch in ctrl_chunks["pdf"]),
       {k: (ctrl_docs[k]["status"], ctrl_docs[k].get("error_message"), (v or [""])[0][:80]) for k, v in ctrl_chunks.items()})

    # ---- chunk contents
    chunks = all_chunks(main_id, EB)
    text = "\n".join(chunks)
    bad = [ch[:80] for ch in chunks if not SECTION_RE.match(ch)]
    ok("mọi đoạn bảng tính có tiền tố [Sheet … › dòng …]", chunks and not bad, bad[:3])
    leaked = [s for s in FACTS["hidden"] if s in text]
    ok("sheet / dòng / cột ẩn không lọt vào đoạn nào", not leaked, leaked)
    want = ["01/03/2026", "0,02%", "73.412", "＜500 triệu"]
    missing = [w for w in want if w not in text]
    ok("ngày, phần trăm, số tiền, dấu ＜ hiển thị đúng; không còn '<' thô", not missing and "<" not in text,
       f"thiếu {missing}, có '<': {'<' in text}")
    group, n = next(iter(FACTS["group_rows"].items()))
    lines = [ln for ch in chunks for ln in ch.split("\n") if "Dịch vụ: Dịch vụ G1-" in ln]
    ok(f"ô gộp dọc 'Nhóm' lặp trên cả {n} dòng của nhóm",
       len(lines) == n and all(f"Nhóm: {group}" in ln for ln in lines), (len(lines), lines[:2]))

    s = c.get("/v1/usage/summary", headers=A, params={"days": 1, "component": "document_embedding"}).json()
    today = s["by_day"][-1]["total_tokens"] if s.get("by_day") else 0
    ok("thống kê token: nhúng tài liệu hôm nay > 0", today > 0, s.get("by_day"))

    # ---- files that must fail cleanly, never stuck in "indexing"
    bad_docs = {k: upload(*FILES[k]).json() for k in ("empty", "fake")}
    bad_ready = {k: wait_ready(v["id"], EB) for k, v in bad_docs.items()}
    e = bad_ready["empty"]
    ok("bảng tính rỗng: failed, báo 'không có dữ liệu'",
       e["status"] == "failed" and "không có dữ liệu" in (e.get("error_message") or ""), (e["status"], e.get("error_message")))
    f = bad_ready["fake"]
    ok("tệp Word đổi đuôi .xlsx: failed, báo '.xlsx hợp lệ'",
       f["status"] == "failed" and ".xlsx hợp lệ" in (f.get("error_message") or ""), (f["status"], f.get("error_message")))

    # ---- retrieval
    hits = retrieved(QUESTION, top_k=3)
    ok("/v1/retrieval trả về bảng tính", main_id in ids(hits), [(h["filename"], h["score"]) for h in hits])
    ok("kết quả đầu tiên là đoạn '[Sheet Biểu phí'", hits and hits[0]["content"].startswith("[Sheet Biểu phí"),
       hits[0]["content"][:120] if hits else hits)

    # ---- one chat round through a temporary assistant
    if SKIP_CHAT:
        print("SKIP trợ lý trích dẫn sheet / trả lời số liệu (SKIP_CHAT=1)")
    else:
        app = c.post("/v1/apps", headers=EB, json={"name": f"Trợ lý kiểm thử xlsx {tag}", "dataset_ids": [ds["id"]]}).json()
        with c.stream("POST", f"/v1/apps/{app['id']}/test-chat", headers=EB, json={"message": QUESTION}, timeout=300) as r:
            ev = sse(r)
        src = [s for e in ev if e["type"] == "sources" for s in e["sources"]]
        answer = "".join(e["content"] for e in ev if e["type"] == "token")
        ok("trợ lý trích dẫn bảng tính, mục 'Sheet …'",
           any(s.get("document_id") == main_id and (s.get("section") or "").startswith("Sheet ") for s in src),
           [(s.get("filename"), s.get("section")) for s in src])
        ok("câu trả lời có số tiền 73.412", any(a in answer for a in ("73.412", "73412", "73,412")), answer[:200])

    # ---- re-index: same chunks
    r = c.post(f"/v1/documents/{main_id}/index", headers=EB)
    d = wait_ready(main_id, EB) if r.status_code == 202 else {"status": r.status_code}
    again = all_chunks(main_id, EB) if d["status"] == "ready" else []
    ok("lập chỉ mục lại: 202, cùng số đoạn, nội dung y hệt",
       r.status_code == 202 and d.get("chunk_count") == main_count and again == chunks,
       (r.status_code, d.get("status"), d.get("chunk_count"), main_count, len(again)))

    # ---- disable / enable / delete
    c.patch(f"/v1/documents/{main_id}", headers=EB, json={"enabled": False})
    ok("đã tắt: truy vấn không còn bảng tính, .txt đối chứng vẫn ra",
       main_id not in ids(retrieved(QUESTION)) and ctrl_ids["txt"] in ids(retrieved(TXT_QUESTION)))
    c.patch(f"/v1/documents/{main_id}", headers=EB, json={"enabled": True})
    ok("bật lại: truy vấn lại ra bảng tính", main_id in ids(retrieved(QUESTION)))
    c.delete(f"/v1/documents/{main_id}", headers=EB)
    ok("xoá: GET 404 và truy vấn không còn bảng tính",
       c.get(f"/v1/documents/{main_id}", headers=EB).status_code == 404 and main_id not in ids(retrieved(QUESTION)))
finally:
    if app and app.get("id"):
        c.delete(f"/v1/apps/{app['id']}", headers=EB)
    c.delete(f"/v1/datasets/{ds['id']}", headers=EB)
    ok("dọn dẹp kho + trợ lý tạm", c.get(f"/v1/datasets/{ds['id']}", headers=EB).status_code == 404)

passed = sum(1 for _, v in results if v)
print(f"\n===== XLSX KB: {passed}/{len(results)} passed =====")
sys.exit(0 if passed == len(results) else 1)
